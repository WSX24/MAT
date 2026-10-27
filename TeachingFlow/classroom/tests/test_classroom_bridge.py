"""bridge（任意路径 PPTX → 课堂备课）的测试。

全部走 action="ingest"：备课要模型、要网络、要花钱，而这一层该负责的是"把路径接进
既有作业链、如实回报结果" —— 那件事不调模型也能完整验证。

两条纪律：
  - 现场用 python-pptx 生成 3 页小课件，不依赖仓库里那份 42 页真实课件；
  - 把 decks/ 挪到 tmp（见了 decks 夹具）：测试绝不往仓库的 decks/ 里写产物。
"""

from __future__ import annotations

import asyncio
import sys
from dataclasses import replace
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import bridge                                   # noqa: E402
from config import ProviderProfile, load_settings   # noqa: E402
from deck import store as deck_store            # noqa: E402


def make_pptx(path: Path, slides: int = 3) -> Path:
    """现场生成一份小 pptx：3 页足够，但少于 3 页就测不出"多页都解析了"。"""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i in range(slides):
        page = prs.slides.add_slide(prs.slide_layouts[5])       # Title Only
        page.shapes.title.text = f"第 {i + 1} 页：课堂测试"
        box = page.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2))
        box.text_frame.text = f"第 {i + 1} 页正文，让解析器有东西可读。"
    prs.save(str(path))
    return path


@pytest.fixture()
def decks(tmp_path, monkeypatch) -> Path:
    """把课件仓库挪到 tmp —— deck.json 落盘位置与 plan 统计都跟着它走。"""
    root = tmp_path / "decks"
    root.mkdir()
    monkeypatch.setattr(deck_store, "DECKS_DIR", root)
    return root


def _dead_settings():
    """一份"连不上"的配置：127.0.0.1:1 没人监听，探活立刻 refused，一个网络请求都不发。"""
    base = load_settings()
    dead = ProviderProfile(name="ollama", base_url="http://127.0.0.1:1/v1",
                           api_key="ollama", model="dead-model")
    return replace(base, provider="ollama", providers={"ollama": dead})


# ---------------------------------------------------------------------------
# 正常路径
# ---------------------------------------------------------------------------

def test_ingest_writes_deck_json(tmp_path, decks) -> None:
    """任意路径的 pptx → decks/<deck_id>/deck.json 落盘，且不复制源文件。"""
    src = make_pptx(tmp_path / "第1单元-测试课件.pptx")
    res = asyncio.run(bridge.preprocess_pptx(src, action="ingest"))

    assert res["ok"] is True, res["error"]
    deck_id = deck_store.make_deck_id(src.name)
    assert res["deck_id"] == deck_id, res["deck_id"]
    assert res["source"] == str(src)
    assert res["action"] == "ingest"
    assert not res["error"]
    assert (decks / deck_id / "deck.json").is_file(), "deck.json 没落盘"
    assert res["stats"]["slides"] == 3, res["stats"]
    assert res["stats"]["has_deck"] is True
    assert "ingest" in res["stages"], res["stages"]
    assert res["degraded"] is False
    # 源文件还在原处（没有被搬走，也没有被拷进 pptx/）
    assert src.is_file()


def test_repeated_calls_land_on_the_same_deck(tmp_path, decks) -> None:
    """同一份源再来一次：不许炸，deck_id 也不许越攒越多。"""
    src = make_pptx(tmp_path / "重复解析.pptx")
    first = asyncio.run(bridge.preprocess_pptx(src, action="ingest"))
    second = asyncio.run(bridge.preprocess_pptx(src, action="ingest"))
    assert first["ok"] is True and second["ok"] is True, (first["error"], second["error"])
    assert first["deck_id"] == second["deck_id"]
    assert len([p for p in decks.iterdir() if p.is_dir()]) == 1


async def test_concurrent_calls_keep_single_flight(tmp_path, decks) -> None:
    """并发两次：单飞语义照旧 —— 一份跑、一份被明确拒掉，都不许抛。"""
    src = make_pptx(tmp_path / "并发课件.pptx")
    first, second = await asyncio.gather(
        bridge.preprocess_pptx(src, action="ingest"),
        bridge.preprocess_pptx(src, action="ingest"),
    )
    assert first["ok"] is True, first["error"]
    assert second["ok"] is False, "同一份源同时在跑，第二份必须被拒（否则会写出半截 plan.json）"
    assert "正在预处理" in second["error"], second["error"]
    # 被拒的那份也是同形的返回值，不是异常
    assert set(second) >= {"ok", "deck_id", "source", "action", "error", "stages", "stats"}


# ---------------------------------------------------------------------------
# 不抛异常（bridge 的硬约束）
# ---------------------------------------------------------------------------

def test_missing_source_is_reported_not_raised(tmp_path, decks) -> None:
    res = asyncio.run(bridge.preprocess_pptx(tmp_path / "没有这份.pptx", action="ingest"))
    assert res["ok"] is False
    assert "不存在" in res["error"], res["error"]


def test_non_pptx_is_rejected(tmp_path, decks) -> None:
    f = tmp_path / "讲义.docx"
    f.write_bytes(b"not a pptx")
    res = asyncio.run(bridge.preprocess_pptx(f, action="ingest"))
    assert res["ok"] is False
    assert ".pptx" in res["error"], res["error"]


def test_unknown_action_is_rejected(tmp_path, decks) -> None:
    src = make_pptx(tmp_path / "小课件.pptx")
    res = asyncio.run(bridge.preprocess_pptx(src, action="cook"))
    assert res["ok"] is False
    assert "action" in res["error"], res["error"]


# ---------------------------------------------------------------------------
# 同步包装（flow.py 走的就是这条）
# ---------------------------------------------------------------------------

def test_sync_wrapper_without_running_loop(tmp_path, decks) -> None:
    src = make_pptx(tmp_path / "同步包装.pptx")
    res = bridge.preprocess_pptx_sync(src, action="ingest")
    assert res["ok"] is True, res["error"]


async def test_sync_wrapper_inside_running_loop(tmp_path, decks) -> None:
    """已经有事件循环时（uvicorn 里）必须能跑 —— 这正是"新线程 + 新 loop"存在的理由。"""
    src = make_pptx(tmp_path / "循环里调同步包装.pptx")
    res = bridge.preprocess_pptx_sync(src, action="ingest")
    assert res["ok"] is True, res["error"]


# ---------------------------------------------------------------------------
# 没有可用模型时的降级（不发任何网络请求）
# ---------------------------------------------------------------------------

def test_model_ready_flags_a_dead_local_service() -> None:
    ok, why = bridge.model_ready(_dead_settings())
    assert ok is False
    assert "连不上" in why, why


def test_model_ready_needs_a_key_for_remote_providers() -> None:
    base = load_settings()
    remote = ProviderProfile(name="deepseek", base_url="https://api.deepseek.com/v1",
                             api_key="", model="deepseek-chat")
    ok, why = bridge.model_ready(
        replace(base, provider="deepseek", providers={"deepseek": remote}))
    assert ok is False
    assert "api_key" in why, why


def test_full_without_model_falls_back_to_rule_plan(tmp_path, decks) -> None:
    """full 但没模型：不碰模型，直接规则兜底 —— 与直连入口（run.py plan --rule）同一口径。

    这条以前钉的是"降级成只解析"：同一份 pptx 走出片链路只有 deck.json、走网页却有
    一份规则计划，两条入口给出两份不同的产物（P2-2）。现在两处都落到规则兜底，
    所以断言翻过来：plan.json 要在，而且 clarify 原因的 note 照样在。
    """
    src = make_pptx(tmp_path / "没模型也能备课.pptx")
    res = asyncio.run(bridge.preprocess_pptx(src, action="full", settings=_dead_settings()))
    assert res["degraded"] is True, res["note"]
    assert res["action"] == "full", "规则兜底是 full 里的一段，不是另一档 action"
    assert res["ok"] is True, res["error"]
    assert res["stats"]["has_deck"] is True
    assert res["stats"]["has_plan"] is True, "无模型时也要落 plan.json"
    assert "plan" in res["stages"], "备课那一段照跑，只是不调模型"
    assert "未配置可用模型" in res["note"], res["note"]
    assert "规则兜底" in res["note"], res["note"]


def test_plan_without_model_is_an_honest_failure(tmp_path, decks) -> None:
    """action=plan 本身就是模型那一段活儿：没模型就如实失败，不假装成功。"""
    src = make_pptx(tmp_path / "只备课没模型.pptx")
    res = asyncio.run(bridge.preprocess_pptx(src, action="plan", settings=_dead_settings()))
    assert res["ok"] is False
    assert "未配置可用模型" in res["error"], res["error"]
