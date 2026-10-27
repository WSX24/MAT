# -*- coding: utf-8 -*-
"""studio/catalog.py：扫描口径（哪些目录算课程、id 稳不稳、坏输入怎么说话）。

为什么这些用例值得存在：课程库的全部事实来自盘上的目录布局，扫错一次的表现是
"学生看不到课"或者"看到一排点进去空白的卡片"；而 id 不稳定会让每次「重新扫描」
把同一门课重插一遍，进度跟着丢。所以这里钉的是行为，不是实现细节。

造交付件目录名一律走 core/naming.py —— 它是命名口径的唯一出处，测试里自己拼
字符串就等于又立一套口径（那样测试会替实现背书，而不是替口径背书）。
"""

from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from studio import catalog                  # noqa: E402


def _naming():
    """core/naming.py：按文件加载，不往 sys.path 里塞 TeachingFlow/。

    为什么这么小心：那份 naming.py 就在 TeachingFlow/core/ 下，而 classroom/ 是
    另一个包根。为一条 import 把 TeachingFlow/ 加进全局 sys.path，会让本文件的
    副作用影响到同一个 pytest 进程里的其它用例（那边也有 tools、fixtures 这类名字）。
    """
    path = Path(__file__).resolve().parents[2] / "core" / "naming.py"
    spec = importlib.util.spec_from_file_location("_test_core_naming", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


NAMING = _naming()


# ===========================================================================
# 夹具：现场造交付件（不碰仓库里的任何真目录）
# ===========================================================================

def _pptx(path: Path, slides: int = 2) -> Path:
    """造一份小课件（只有标题 + 一行正文，解析器认得出、不用联网）。"""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i in range(slides):
        page = prs.slides.add_slide(prs.slide_layouts[5])       # Title Only
        page.shapes.title.text = f"第 {i + 1} 页"
        box = page.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1.5))
        box.text_frame.text = f"第 {i + 1} 页正文，讲两句就够。"
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def _unit_path(course: Path, ch_no: int, ch_title: str, u_no: int, u_title: str) -> Path:
    """交付件路径：章目录只由 naming 造，文件名只由 naming 造。"""
    ch = NAMING.chapter_dir(course, {"chapter_order": ch_no, "chapter_title": ch_title})
    label = NAMING.unit_label({"unit_index": u_no - 1, "unit_title": u_title})
    return ch / f"{label}.pptx"


@pytest.fixture()
def runs(tmp_path) -> Path:
    """runs 根目录（还不存在，由各用例自己往里放课程目录）。"""
    return tmp_path / "runs"


def _write_run(runs: Path, *, name: str = "我想系统学习高中物理_144652",
               date: str | None = "20261027",
               chapters=((1, "力学入门", [(1, "运动的描述"), (2, "速度与加速度")]),
                         (2, "电学入门", [(1, "电荷与电流")])),
               manifest: dict | None = None) -> Path:
    """造一个 run：2 章 3 单元，外加可选 manifest.json。"""
    course = (runs / date / name) if date else (runs / name)
    course.mkdir(parents=True, exist_ok=True)
    for ch_no, ch_title, units in chapters:
        for u_no, u_title in units:
            _pptx(_unit_path(course, ch_no, ch_title, u_no, u_title), slides=2)
    if manifest is not None:
        (course / "manifest.json").write_text(
            json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    return course


# ===========================================================================
# 用例
# ===========================================================================

def test_scan_skips_a_mock_run_without_pptx(runs: Path) -> None:
    """mock 运行没有交付件 → 不产出课程、计数被记下来、说清为什么跳过。"""
    course = runs / "20261027" / "我想简单学一下《祥林嫂》这篇课文_143641"
    course.mkdir(parents=True)
    (course / "manifest.json").write_text(
        json.dumps({"input": "我想简单学一下《祥林嫂》这篇课文"}, ensure_ascii=False),
        encoding="utf-8")

    res = catalog.scan_runs(runs)

    assert res == [], "没有 pptx 的目录不该变成课程"
    assert res.ignored == 1
    assert any("已跳过" in w for w in res.warnings), res.warnings
    assert any("143641" in w for w in res.warnings), "警告里要说清是哪一门被跳过了"


def test_scan_reads_two_chapters_three_units(runs: Path) -> None:
    """2 章 3 单元 → 1 门课 3 个单元，章序/单元序都折进 order_index。"""
    _write_run(runs)

    res = catalog.scan_runs(runs)

    assert len(res) == 1
    course = res[0]
    assert course.title == "我想系统学习高中物理", "标题取目录名去掉 _时分秒"
    assert course.source == "run"
    assert course.units and all(u.id and u.pptx_path for u in course.units)
    assert [u.title for u in course.units] == ["运动的描述", "速度与加速度", "电荷与电流"]
    assert [u.chapter for u in course.units] == ["力学入门", "力学入门", "电学入门"]
    assert [u.order_index for u in course.units] == [0, 1, 2], "全课连续，排序即章序"
    assert [u.slides for u in course.units] == [2, 2, 2], "页数扫描时就数出来"
    assert all(u.est_sec is None for u in course.units), "没解析过不猜时长"
    assert all(Path(u.pptx_path).is_file() for u in course.units)


def test_scan_ids_are_stable_across_scans(runs: Path) -> None:
    """同一输入扫两次，课程与单元的 id 一个字符都不变（upsert 的前提）。"""
    _write_run(runs)

    first = catalog.scan_runs(runs)
    second = catalog.scan_runs(runs)

    assert [c.id for c in first] == [c.id for c in second]
    assert [u.id for u in first[0].units] == [u.id for u in second[0].units]
    assert len({u.id for u in first[0].units}) == 3, "同一门课的单元 id 不许互撞"
    assert all(len(u.id) == 12 and len(first[0].id) == 12 for u in first[0].units)


def test_manifest_supplies_subject_and_title(runs: Path) -> None:
    """manifest 可读时用它的 input/spec.subject 补标题与科目。"""
    _write_run(runs, manifest={"input": "我想系统学习高中物理", "spec": {"subject": "物理"}})

    course = catalog.scan_runs(runs)[0]

    assert course.title == "我想系统学习高中物理"
    assert course.subject == "物理"


def test_old_manifest_subject_in_args_is_also_read(runs: Path) -> None:
    """老 run 的科目只在 args 里 —— 两种取处都要认。"""
    _write_run(runs, manifest={"input": "高一物理第一课", "args": {"subject": "物理"}})

    course = catalog.scan_runs(runs)[0]

    assert course.subject == "物理"


def test_broken_manifest_falls_back_to_dir_title(runs: Path) -> None:
    """manifest 是坏 JSON：课程照常出现，标题退回目录名，警告里说清。"""
    course = _write_run(runs, name="坏manifest_101010")
    (course / "manifest.json").write_text("{ 这不是 json", encoding="utf-8")

    res = catalog.scan_runs(runs)

    assert len(res) == 1
    assert res[0].title == "坏manifest"
    assert res[0].subject == ""
    assert any("manifest.json 读不动" in w for w in res.warnings), res.warnings


def test_missing_runs_root_is_empty_and_says_so(tmp_path: Path) -> None:
    """runs 根目录不存在：空列表 + 一句人话，不抛。"""
    res = catalog.scan_runs(tmp_path / "从来没有过" / "runs")

    assert res == []
    assert res.ignored == 0
    assert any("runs 根目录不存在" in w for w in res.warnings), res.warnings


def test_course_dir_without_date_layer_is_also_found(runs: Path) -> None:
    """更早的布局（runs/<课程>/，没有日期层）也要认出来。"""
    _write_run(runs, date=None)

    res = catalog.scan_runs(runs)

    assert len(res) == 1
    assert res[0].title == "我想系统学习高中物理"


def test_stray_pptx_outside_chapter_layout_is_reported_not_counted(runs: Path) -> None:
    """旧版布局的 pptx（05_ppt_xxx/ 下）：不算单元，但要在警告里说一句。"""
    course = _write_run(runs)
    old = course / "05_ppt_ch01-u1"
    _pptx(old / "旧产物.pptx", slides=1)

    res = catalog.scan_runs(runs)

    assert len(res[0].units) == 3, "旧产物不该被算成第 4 个单元"
    assert any("旧版布局" in w and "已忽略" in w for w in res.warnings), res.warnings


def test_unit_from_pptx_counts_slides_and_tolerates_broken_file(tmp_path: Path) -> None:
    """unit_from_pptx：好文件数得出页数，坏文件给 None（扫描不该因为一份坏课件崩）。"""
    good = _pptx(tmp_path / "第1单元-好课件.pptx", slides=3)

    unit = catalog.unit_from_pptx(good, "course1", "好课件", "力学入门", 4)

    assert unit.slides == 3
    assert unit.order_index == 4
    assert unit.chapter == "力学入门"
    assert unit.est_sec is None

    bad = tmp_path / "第2单元-坏课件.pptx"
    bad.write_text("这不是 pptx", encoding="utf-8")
    broken = catalog.unit_from_pptx(bad, "course1", "坏课件", "力学入门", 5)

    assert broken.slides is None
    assert broken.id != unit.id
