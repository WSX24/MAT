"""API 契约测试（设计文档 §4 逐条钉住）。

为什么这些用例必须存在：前端**只读契约**。后端把 position.slide 改成 position_slide
不会报错，只会让 W2 的页面在某处悄悄不显示 —— 这类回归只有"逐字段断言"能抓住。

几条测试纪律：

* **每例一个临时库**（tmp_path + CLASSROOM_DB）与临时 STATE_ROOT：
  跑完不留痕，也不受"上一次跑过什么"影响；
* **不联网、不调模型、不写仓库目录**：W3 的内容管线入口与 TTS 全部替身化，
  所以这些用例在没有 key、没有网的机器上必须全绿；
* 既要断言"形状"也要断言"取值"：只断言 200 的用例抓不到字段缺失。

WS 用例走**真运行器 + stub 课件**（不是假运行器）：连接、暂停、续播、翻页这些
真正会出问题的接线，只有跑真代码才验得到。
"""

from __future__ import annotations

import asyncio
import json
import queue as queue_mod
import sys
import threading
import time
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

CLASSROOM = Path(__file__).resolve().parents[1]
if str(CLASSROOM) not in sys.path:              # 允许从任意 cwd 直接跑本文件
    sys.path.insert(0, str(CLASSROOM))

from store import db as db_mod                  # noqa: E402
from store import repo                          # noqa: E402

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


# ===========================================================================
# 夹具与替身
# ===========================================================================

@pytest.fixture(autouse=True)
def _isolate(tmp_path, monkeypatch):
    """每个用例：独立库、独立 imports 目录、独立 dist、/dev 关。"""
    monkeypatch.setenv("CLASSROOM_DB", str(tmp_path / "classroom.sqlite"))
    monkeypatch.delenv("CLASSROOM_DEV", raising=False)
    from app import deps

    monkeypatch.setattr(deps, "STATE_ROOT", tmp_path / "state", raising=False)
    monkeypatch.setattr(deps, "DIST_ROOT", tmp_path / "dist", raising=False)
    deps.reset_settings_cache()
    yield
    deps.reset_settings_cache()


@pytest.fixture()
def client(_isolate):
    from app import create_app

    with TestClient(create_app()) as c:
        yield c


def _conn():
    """直连测试库。`_conn()` 也跑迁移：app 的迁移是懒的（第一条请求才建表），
    而有些用例要先于任何请求直接写库。"""
    conn = db_mod.connect()
    db_mod.migrate(conn)
    return conn


def _seed_course(units: int = 2, *, course_id: str = "c1", with_pptx: bool = True) -> list[str]:
    conn = _conn()
    try:
        repo.upsert_course(conn, {"id": course_id, "title": "测试课程", "subject": "物理",
                                  "source": "run", "run_dir": "/tmp/run1"})
        ids = []
        for i in range(units):
            uid = f"{course_id}-u{i + 1}"
            repo.upsert_unit(conn, {
                "id": uid, "course_id": course_id, "chapter": "第1章",
                "title": f"单元{i + 1}", "order_index": i,
                "pptx_path": f"/tmp/{course_id}/u{i + 1}.pptx" if with_pptx else "",
                "slides": 3, "est_sec": 60.0,
            })
            ids.append(uid)
        return ids
    finally:
        conn.close()


def _fake_prepare(unit_row: dict) -> dict:
    return {"ok": True, "deck_ready": True, "plan_ready": True, "plan_source": "rule",
            "degraded": True, "warnings": ["用的是规则兜底"], "slides": 7, "est_sec": 123.0}


def _fake_outline(unit_row: dict) -> dict:
    return {"totals": {"slides": 3, "units": 2, "est_sec": 60.0},
            "pages": [{"index": 0, "title": "第1页",
                       "units": [{"preview": "讲这个", "est_sec": 10.0, "kp": ["力"]}]}]}


def _fake_scan() -> tuple[list[dict], list[str]]:
    return ([{
        "id": "run-1", "title": "扫描来的课", "subject": "物理", "source": "run",
        "run_dir": "/tmp/runs/20260101/扫描来的课_101010",
        "units": [
            {"id": "run-1-u1", "title": "第1单元", "chapter": "第1章", "order_index": 0,
             "pptx_path": "/tmp/runs/第1单元.pptx", "slides": 4, "est_sec": None},
            {"id": "run-1-u2", "title": "第2单元", "chapter": "第1章", "order_index": 1,
             "pptx_path": "/tmp/runs/第2单元.pptx", "slides": 6, "est_sec": None},
        ],
    }], [])


def _make_pptx(path: Path, slides: int = 2) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    for i in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(0, 0, 200, 80)
        box.text_frame.text = f"第{i + 1}页"
    prs.save(path)
    return path.read_bytes()


# ---- 运行器替身（WS 用例用真运行器 + stub 课件） --------------------------

def _mini_settings():
    from config import ProviderProfile, Settings

    return Settings(
        provider="stub",
        providers={"stub": ProviderProfile(name="stub", base_url="http://127.0.0.1:1/v1",
                                           api_key="", model="stub")},
        temperature=0.2, timeout=5.0,
        tts_enabled=False, tts_voice="", tts_rate="+0%",
        probe_window=0.05, max_clarify_asks=1, max_rephrase=1, unit_char_warn=200,
    )


class _StubLLM:
    """报告"没有模型"的替身：讲解路径根本不该碰它。"""

    degraded = True
    down_reason = "测试替身"

    def mark_down(self, *_a, **_k):
        return None

    def mark_up(self):
        return None

    def think_for(self, _role: str) -> bool:
        return False


def _stub_course(slides: int = 2, units_per_slide: int = 2):
    from course import Course
    from models import LessonScript, Slide, SlideElement
    from state import LectureUnit

    slide_list, units = [], []
    for i in range(slides):
        slide_list.append(Slide(index=i, title=f"第{i + 1}页", elements=[
            SlideElement(element_id=f"s{i + 1}-t1", kind="text", text="内容", order=0)]))
        for j in range(units_per_slide):
            units.append(LectureUnit(
                unit_id=f"u{i + 1}-{j + 1}", slide=i, index_in_slide=j,
                text=f"第{i + 1}页第{j + 1}段。第二句话。",
                is_page_end=(j == units_per_slide - 1), kp=[]))
    script = LessonScript(course_id="stub", title="假课件", session_goals=[],
                          slides=slide_list, units=units)
    return Course(script=script, mode="simple")


def _install_stub_engine(monkeypatch, *, slides: int = 4, speech_timeout: float = 0.3):
    from app import ws as ws_mod
    from studio.runner import SessionRunner

    def fake_build_runner(*, session, unit, profile, send, conn):
        return SessionRunner(
            session=session, unit=unit, profile=profile, send=send, conn=conn,
            settings=_mini_settings(), tts=None, course=_stub_course(slides),
            llm=_StubLLM(), speech_timeout=speech_timeout)

    monkeypatch.setattr(ws_mod, "build_runner", fake_build_runner)


class _Collector:
    def __init__(self):
        self.frames: list[dict] = []

    async def __call__(self, kind: str, payload: dict) -> None:
        self.frames.append({"kind": kind, **payload})

    def kinds(self) -> list[str]:
        return [f["kind"] for f in self.frames]


async def _await_until(predicate, timeout: float = 3.0) -> bool:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if predicate():
            return True
        await asyncio.sleep(0.01)
    return False


def _session_via_api(client, unit_id: str) -> dict:
    resp = client.post("/api/v1/sessions", json={"unit_id": unit_id})
    assert resp.status_code == 200, resp.text
    return resp.json()


def _pump(ws, *, until, limit: int = 60) -> list[dict]:
    """读到满足 until 的那一帧（或读满 limit）。"""
    frames = []
    for _ in range(limit):
        frame = ws.receive_json()
        frames.append(frame)
        if until(frame):
            return frames
    return frames


def _drain_until_state(ws, marker: str, *, limit: int = 40) -> list[dict]:
    """读到「某条 notyfy 之后的那个 state 帧」。

    TestClient 的 WS 通道是**会合式**（buffer 为 0）：服务端多推一帧、而测试这边
    不读，双方就各等各的 —— 表现为用例挂死。暂停/继续都是"先 notify 再 state"
    两帧连发，所以断言点必须落在第二帧上，把两帧都读掉。
    """
    frames: list[dict] = []
    seen_marker = False
    for _ in range(limit):
        frame = ws.receive_json()
        frames.append(frame)
        if frame["kind"] == "notify" and marker in (frame.get("text") or ""):
            seen_marker = True
            continue
        if seen_marker and frame["kind"] == "state":
            return frames
    return frames


def _frame_reader(ws) -> "queue_mod.Queue":
    """把 WS 帧搬进一个队列，测试侧就能"带超时地等一帧 / 断言 N 秒内没有帧"。

    为什么需要它：TestClient 的 WS 通道是**会合式**（buffer 0），`receive_json()`
    没有超时参数 —— 想断言"这段时间不该有任何帧"就只能靠一条后台线程持续读，
    再用 `queue.get(timeout=...)` 判断。顺带它也保证了服务端不会被"没人读"卡住。
    """
    frames: "queue_mod.Queue" = queue_mod.Queue()

    def _loop() -> None:
        try:
            while True:
                frames.put(ws.receive_json())
        except Exception:          # 连接关闭 / AssertionError 之类：读线程安静退出
            return

    threading.Thread(target=_loop, name="ws-frame-reader", daemon=True).start()
    return frames


def _next_frame(frames, timeout: float = 5.0) -> dict | None:
    try:
        return frames.get(timeout=timeout)
    except queue_mod.Empty:
        return None


def _wait_kind(frames, kind: str, *, timeout: float = 5.0) -> dict | None:
    deadline = time.monotonic() + timeout
    while True:
        left = deadline - time.monotonic()
        if left <= 0:
            return None
        frame = _next_frame(frames, left)
        if frame is None:
            return None
        if frame.get("kind") == kind:
            return frame


# ===========================================================================
# 健康与静态
# ===========================================================================

def test_health_always_ok(client):
    r = client.get("/api/health")
    assert r.status_code == 200
    body = r.json()
    assert body["ok"] is True
    assert body["version"]
    assert body["courses"] == 0
    assert set(body["engine"]) == {"provider", "model", "tts"}


def test_root_returns_200_without_dist(client):
    """dist 不存在也要 200：并行开发时后端必须能起。"""
    r = client.get("/")
    assert r.status_code == 200
    assert "text/html" in r.headers["content-type"]
    assert "后端已经起来了" in r.text


def test_root_serves_dist_index_when_built(client, monkeypatch, tmp_path):
    from app import deps

    dist = tmp_path / "dist"
    (dist / "assets").mkdir(parents=True)
    (dist / "index.html").write_text("<!doctype html><title>课堂</title>REAL_APP", encoding="utf-8")
    (dist / "assets" / "app.js").write_text("console.log(1)", encoding="utf-8")
    monkeypatch.setattr(deps, "DIST_ROOT", dist, raising=False)
    assert "REAL_APP" in client.get("/").text
    assert client.get("/assets/app.js").text == "console.log(1)"
    assert client.get("/assets/missing.js").status_code == 404
    # SPA 深链刷新也要能打开
    assert "REAL_APP" in client.get("/learn/u1").text


def test_unknown_api_is_404_not_spa(client):
    r = client.get("/api/v1/nope")
    assert r.status_code == 404
    assert r.json()["ok"] is False
    assert "找不到" in r.json()["message"]
    # 没有尾斜杠的 /api 同样不能落到 SPA 首页（否则排查时会看到"接口返回 HTML"）
    assert client.get("/api").status_code == 404


# ===========================================================================
# 档案
# ===========================================================================

def test_profiles_default_and_shape(client):
    r = client.get("/api/v1/profiles")
    assert r.status_code == 200
    body = r.json()
    assert body["active_id"]
    assert len(body["items"]) == 1
    assert set(body["items"][0]) == {"id", "name", "created_at"}


def test_profile_create_and_activate(client):
    created = client.post("/api/v1/profiles", json={"name": "小明"})
    assert created.status_code == 200
    new_id = created.json()["id"]
    assert created.json()["name"] == "小明"
    assert client.post(f"/api/v1/profiles/{new_id}/activate").json() == {"ok": True}
    listed = client.get("/api/v1/profiles").json()
    assert listed["active_id"] == new_id
    assert [p["name"] for p in listed["items"]] == ["默认档案", "小明"]


def test_profile_rejects_empty_name(client):
    r = client.post("/api/v1/profiles", json={"name": "   "})
    assert r.status_code == 400
    assert "档案名不能为空" in r.json()["message"]


def test_profile_delete_cascades(client):
    created = client.post("/api/v1/profiles", json={"name": "临时"}).json()
    pid = created["id"]
    unit_ids = _seed_course(1)
    sess = client.post("/api/v1/sessions", json={"unit_id": unit_ids[0],
                                                 "profile_id": pid}).json()
    conn = _conn()
    try:
        repo.add_event(conn, sess["session_id"], "ping", {})
        repo.upsert_progress(conn, pid, unit_ids[0], percent=12.0)
    finally:
        conn.close()
    assert client.delete(f"/api/v1/profiles/{pid}").json() == {"ok": True}
    assert client.delete(f"/api/v1/profiles/{pid}").status_code == 404
    conn = _conn()
    try:
        assert conn.execute("SELECT COUNT(*) AS n FROM sessions WHERE profile_id=?",
                            (pid,)).fetchone()["n"] == 0
        assert conn.execute("SELECT COUNT(*) AS n FROM events WHERE session_id=?",
                            (sess["session_id"],)).fetchone()["n"] == 0
        assert conn.execute("SELECT COUNT(*) AS n FROM progress WHERE profile_id=?",
                            (pid,)).fetchone()["n"] == 0
    finally:
        conn.close()


# ===========================================================================
# 课程
# ===========================================================================

def test_courses_empty(client):
    assert client.get("/api/v1/courses").json() == []


def test_courses_scan_upserts_and_counts(client, monkeypatch):
    from app import deps

    monkeypatch.setattr(deps, "scan_runs", _fake_scan)
    first = client.post("/api/v1/courses/scan").json()
    assert (first["added"], first["updated"], first["total"]) == (1, 0, 1)
    second = client.post("/api/v1/courses/scan").json()
    # 同一份输入再扫一次：不是又加一门，而是更新
    assert (second["added"], second["updated"], second["total"]) == (0, 1, 1)
    listed = client.get("/api/v1/courses").json()
    assert len(listed) == 1
    item = listed[0]
    assert item["id"] == "run-1"
    assert item["unit_count"] == 2
    assert item["unit_ready_count"] == 0
    assert set(item["progress"]) == {"started", "completed", "percent"}
    assert set(item) >= {"id", "title", "subject", "source", "unit_count",
                         "unit_ready_count", "progress", "updated_at"}


def test_courses_scan_without_pipeline_degrades(client, monkeypatch):
    from app import deps

    monkeypatch.setattr(deps, "scan_runs",
                        lambda *a, **k: ([], ["内容管线未就绪（studio/catalog.py 缺失）"]))
    body = client.post("/api/v1/courses/scan").json()
    assert body["total"] == 0
    assert body["warnings"]


def test_course_detail_chapters(client):
    unit_ids = _seed_course(2)
    conn = _conn()
    try:
        repo.upsert_unit(conn, {"id": "c1-u3", "course_id": "c1", "chapter": "第2章",
                                "title": "单元3", "order_index": 2, "slides": 5})
    finally:
        conn.close()
    detail = client.get("/api/v1/courses/c1").json()
    assert detail["id"] == "c1" and detail["title"] == "测试课程"
    assert [c["title"] for c in detail["chapters"]] == ["第1章", "第2章"]
    assert [u["id"] for u in detail["chapters"][0]["units"]] == unit_ids
    unit = detail["chapters"][0]["units"][0]
    assert set(unit) == {"id", "title", "slides", "est_sec", "ready", "plan_source", "progress"}
    assert set(unit["progress"]) == {"percent", "last_slide", "status"}
    assert unit["ready"] is False


def test_course_detail_404(client):
    assert client.get("/api/v1/courses/nope").status_code == 404


def test_course_delete_leaves_no_orphans(client):
    """删课程 = 单元、会话、进度、事件一起走。留一行孤儿都会让下次导入"又出现"。"""
    unit_ids = _seed_course(2)
    sess = _session_via_api(client, unit_ids[0])
    conn = _conn()
    try:
        repo.add_event(conn, sess["session_id"], "ping", {"n": 1})
        repo.upsert_progress(conn, client.get("/api/v1/profiles").json()["active_id"],
                             unit_ids[0], percent=42.0, last_slide=2)
    finally:
        conn.close()

    assert client.delete("/api/v1/courses/c1").json() == {"ok": True}
    assert client.get("/api/v1/courses/c1").status_code == 404
    conn = _conn()
    try:
        for sql, arg in (("SELECT COUNT(*) AS n FROM units WHERE course_id=?", "c1"),
                         ("SELECT COUNT(*) AS n FROM sessions WHERE unit_id=?", unit_ids[0]),
                         ("SELECT COUNT(*) AS n FROM progress WHERE unit_id=?", unit_ids[0]),
                         ("SELECT COUNT(*) AS n FROM events WHERE session_id=?", sess["session_id"])):
            assert conn.execute(sql, (arg,)).fetchone()["n"] == 0, sql
    finally:
        conn.close()


def test_courses_list_progress_fields(client):
    unit_ids = _seed_course(2)
    pid = client.get("/api/v1/profiles").json()["active_id"]
    conn = _conn()
    try:
        repo.upsert_progress(conn, pid, unit_ids[0], percent=100.0, last_slide=2)
        repo.upsert_progress(conn, pid, unit_ids[1], percent=30.0, last_slide=1)
    finally:
        conn.close()
    prog = client.get(f"/api/v1/courses?profile_id={pid}").json()[0]["progress"]
    assert prog["started"] == 2
    assert prog["completed"] == 1
    assert prog["percent"] == pytest.approx(65.0)


# ===========================================================================
# 导入
# ===========================================================================

def test_import_rejects_non_pptx(client):
    r = client.post("/api/v1/courses/import",
                    files={"file": ("大纲.txt", b"hello", "text/plain")})
    assert r.status_code == 400
    assert "只支持 .pptx" in r.json()["message"]


def test_import_rejects_broken_pptx(client):
    r = client.post("/api/v1/courses/import",
                    files={"file": ("坏课件.pptx", b"not-a-zip-at-all", PPTX_MIME)})
    assert r.status_code == 400
    assert "不是有效的 PPTX" in r.json()["message"]
    assert client.get("/api/v1/courses").json() == []


def test_import_real_pptx_lands_in_library(client, monkeypatch, tmp_path):
    from app import deps

    def _prepare_without_slides(unit_row):
        out = _fake_prepare(unit_row)
        out["slides"] = None          # 备课没报页数 → 页数取导入时数出来的那份
        return out

    monkeypatch.setattr(deps, "prepare_unit", _prepare_without_slides)
    data = _make_pptx(tmp_path / "物理第一课.pptx", slides=2)
    r = client.post("/api/v1/courses/import",
                    files={"file": ("物理第一课.pptx", data, PPTX_MIME)})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["course_id"].startswith("import-")
    assert body["unit_id"]
    assert body["slides"] == 2
    courses = client.get("/api/v1/courses").json()
    assert [c["source"] for c in courses] == ["import"]
    assert courses[0]["title"] == "物理第一课"
    assert courses[0]["unit_ready_count"] == 1        # 导入时备过一次 → ready
    detail = client.get(f"/api/v1/courses/{body['course_id']}").json()
    assert detail["chapters"][0]["units"][0]["id"] == body["unit_id"]
    # 上传的文件落在隔离的 STATE_ROOT 里（不打进仓库）
    assert (tmp_path / "state" / "imports" / body["course_id"]).is_dir()


# ===========================================================================
# 单元媒体（page 帧 media_base 的落点，W2 已按此实现）
# ===========================================================================

def test_unit_media_serves_deck_image(client, monkeypatch, tmp_path):
    import deck.store as deck_store

    monkeypatch.setattr(deck_store, "DECKS_DIR", tmp_path / "decks", raising=False)
    pptx = tmp_path / "第一课.pptx"
    pptx.write_bytes(b"x")
    course_id = "c-media"
    unit_id = "c-media-u1"
    conn = _conn()
    try:
        repo.upsert_course(conn, {"id": course_id, "title": "媒体课", "source": "import"})
        repo.upsert_unit(conn, {"id": unit_id, "course_id": course_id, "title": "单元",
                               "order_index": 0, "pptx_path": str(pptx), "slides": 1})
    finally:
        conn.close()
    deck_id = deck_store.make_deck_id(pptx.name)
    media = tmp_path / "decks" / deck_id / "media"
    media.mkdir(parents=True)
    # PNG 头 + 一点填充（只要 mime 能识别、内容能原样返回就够）
    (media / "image1.png").write_bytes(bytes.fromhex("89504e470d0a1a0a") + b"0" * 16)

    ok = client.get(f"/api/v1/units/{unit_id}/media/image1.png")
    assert ok.status_code == 200
    assert ok.headers["content-type"].startswith("image/png")
    assert ok.content.startswith(bytes.fromhex("89504e47"))
    # 路径穿越与不存在的文件一律 404（名字来自浏览器，不做"试试看能不能解析"）
    assert client.get(f"/api/v1/units/{unit_id}/media/..%2Fsecret.png").status_code == 404
    assert client.get(f"/api/v1/units/{unit_id}/media/nope.png").status_code == 404
    assert client.get("/api/v1/units/nope/media/image1.png").status_code == 404


def test_session_start_slide(client):
    unit_id = _seed_course(1)[0]
    first = client.post("/api/v1/sessions",
                        json={"unit_id": unit_id, "start_slide": 2}).json()
    assert first["resumed"] is False
    assert first["position"] == {"slide": 2, "unit_index": 0, "from_progress": False}
    # 续学：再要另一页 = 把这条会话挪到那一页（不是另开一条）
    again = client.post("/api/v1/sessions",
                        json={"unit_id": unit_id, "start_slide": 1}).json()
    assert again["session_id"] == first["session_id"]
    assert again["position"] == {"slide": 1, "unit_index": 0, "from_progress": False}
    # 非法值忽略（不报错、不改位置）
    bad = client.post("/api/v1/sessions",
                      json={"unit_id": unit_id, "start_slide": "不是数字"}).json()
    assert bad["session_id"] == first["session_id"]
    assert bad["position"] == {"slide": 1, "unit_index": 0, "from_progress": False}


# ===========================================================================
# 单元
# ===========================================================================

def test_unit_detail_shape(client):
    unit_id = _seed_course(1)[0]
    body = client.get(f"/api/v1/units/{unit_id}").json()
    assert set(body) == {"id", "course_id", "title", "slides", "est_sec",
                         "ready", "plan_source", "progress"}
    assert body["id"] == unit_id and body["slides"] == 3
    assert body["ready"] is False


def test_unit_404(client):
    assert client.get("/api/v1/units/nope").status_code == 404
    assert client.get("/api/v1/units/nope/outline").status_code == 404
    assert client.post("/api/v1/units/nope/prepare").status_code == 404


def test_unit_outline_contract(client, monkeypatch):
    from app import deps

    unit_id = _seed_course(1)[0]
    monkeypatch.setattr(deps, "build_outline", _fake_outline)
    body = client.get(f"/api/v1/units/{unit_id}/outline").json()
    assert set(body) == {"totals", "pages"}
    assert set(body["totals"]) == {"slides", "units", "est_sec"}
    page = body["pages"][0]
    assert set(page) == {"index", "title", "units"}
    assert set(page["units"][0]) == {"preview", "est_sec", "kp"}


def test_unit_outline_degrades_without_pipeline(client, monkeypatch):
    from app import deps

    unit_id = _seed_course(1)[0]
    monkeypatch.setattr(deps, "build_outline",
                        lambda row: {"totals": {"slides": 0, "units": 0, "est_sec": 0.0},
                                     "pages": []})
    body = client.get(f"/api/v1/units/{unit_id}/outline").json()
    assert body["pages"] == []
    assert set(body["totals"]) == {"slides", "units", "est_sec"}


def test_unit_prepare_marks_ready(client, monkeypatch):
    from app import deps

    unit_id = _seed_course(1)[0]
    monkeypatch.setattr(deps, "prepare_unit", _fake_prepare)
    body = client.post(f"/api/v1/units/{unit_id}/prepare").json()
    assert set(body) == {"ok", "deck_ready", "plan_ready", "plan_source",
                         "degraded", "warnings", "slides", "est_sec"}
    assert body["ok"] is True and body["plan_source"] == "rule"
    assert isinstance(body["slides"], int) and body["slides"] == 7
    assert isinstance(body["est_sec"], float) and body["est_sec"] == 123.0
    after = client.get(f"/api/v1/units/{unit_id}").json()
    assert after["ready"] is True                  # 备过就落库
    assert after["plan_source"] == "rule"
    assert after["slides"] == 7                    # 备课结果回填了页数
    assert client.get("/api/v1/courses").json()[0]["unit_ready_count"] == 1


def test_prepare_reuse_path_returns_numbers(client, monkeypatch, tmp_path):
    """F2：备过之后再点一次"准备课件"，页数与时长仍必须是数值。

    复用路径读的是盘上的 deck.json + plan.json（不是本次解析的产物），
    所以这两个数由我们自己按同一口径兜底 —— 上游给 0/None 也不许把契约变成 null，
    否则前端课程卡上就是"— 页 / — 分钟"。
    """
    import deck.store as deck_store
    import studio.materialize as materialize

    monkeypatch.setattr(deck_store, "DECKS_DIR", tmp_path / "decks", raising=False)
    pptx = tmp_path / "复用课件.pptx"
    pptx.write_bytes(b"x")
    deck_id = deck_store.make_deck_id(pptx.name)
    base = tmp_path / "decks" / deck_id
    base.mkdir(parents=True)
    (base / "deck.json").write_text(
        json.dumps({"slides": [{"index": i} for i in range(4)]}), encoding="utf-8")
    # plan.json 的单元按 slide 挂页（缺 slide 的单元挂不上任何一页 → 时长会算成 0，
    # 这正是"两处各算一份"会踩的坑，所以这里按真实格式造）
    (base / "plan.json").write_text(
        json.dumps({"source": "rule", "units": [
            {"slide": 0, "text": "一" * 40},
            {"slide": 1, "text": "二" * 10},
        ]}), encoding="utf-8")

    def fake_prepare(pptx_path, *, deck_root=None, force=False, settings=None):
        # 复刻"复用已有 deck"那一刻：上游只说就绪，两个数值字段是空的
        return {"ok": True, "deck_id": deck_id, "deck_ready": True, "plan_ready": True,
                "plan_source": "rule", "degraded": True, "slides": None, "est_sec": None,
                "warnings": ["已有 deck.json + plan.json，直接复用"]}

    monkeypatch.setattr(materialize, "prepare_unit", fake_prepare)

    unit_id = "c-x-u1"
    conn = _conn()
    try:
        repo.upsert_course(conn, {"id": "c-x", "title": "复用课", "source": "import"})
        repo.upsert_unit(conn, {"id": unit_id, "course_id": "c-x", "title": "复用单元",
                                "order_index": 0, "pptx_path": str(pptx)})
    finally:
        conn.close()

    first = client.post(f"/api/v1/units/{unit_id}/prepare").json()
    assert first["ok"] is True
    assert first["slides"] == 4                       # deck.json 里 4 页
    assert first["est_sec"] == pytest.approx(10.0)    # plan.json：40/5 + 10/5
    second = client.post(f"/api/v1/units/{unit_id}/prepare").json()
    assert second["slides"] == 4 and second["est_sec"] == pytest.approx(10.0)
    # 落库的也是数值（课程卡片读的是这张表）
    assert client.get(f"/api/v1/units/{unit_id}").json()["slides"] == 4


# ===========================================================================
# 会话
# ===========================================================================

def test_session_create_and_get(client):
    unit_id = _seed_course(1)[0]
    created = _session_via_api(client, unit_id)
    assert created["resumed"] is False
    assert created["position"] == {"slide": 0, "unit_index": 0, "from_progress": False}
    sid = created["session_id"]

    body = client.get(f"/api/v1/sessions/{sid}").json()
    assert set(body) == {"id", "unit_id", "status", "started_at", "elapsed_sec",
                         "position", "comprehension", "qa_count"}
    assert body["status"] == "active"
    assert body["position"] == {"slide": 0, "unit_index": 0}
    assert body["comprehension"] == {}
    assert client.get("/api/v1/sessions/nope").status_code == 404


def test_session_resume_returns_same_id(client):
    unit_id = _seed_course(1)[0]
    first = _session_via_api(client, unit_id)
    conn = _conn()
    try:
        repo.update_session(conn, first["session_id"], position_slide=2, position_unit=1)
    finally:
        conn.close()
    again = _session_via_api(client, unit_id)
    assert again["resumed"] is True
    assert again["session_id"] == first["session_id"]
    # 复用"未结束的会话"：位置来自会话本身，不是 progress
    assert again["position"] == {"slide": 2, "unit_index": 1, "from_progress": False}


def test_session_resume_after_end_starts_new(client):
    """结束一条"还没讲任何东西"的会话：新会话仍在第 1 页，resumed=false。"""
    unit_id = _seed_course(1)[0]
    first = _session_via_api(client, unit_id)
    client.post(f"/api/v1/sessions/{first['session_id']}/end")
    second = _session_via_api(client, unit_id)
    assert second["session_id"] != first["session_id"]
    assert second["position"] == {"slide": 0, "unit_index": 0, "from_progress": False}
    assert second["resumed"] is False


def test_session_continues_from_progress_after_end(client):
    """F1：结束一节课后再进同一单元，必须从上次的位置继续（不是第 1 页）。"""
    unit_id = _seed_course(1)[0]
    first = _session_via_api(client, unit_id)
    conn = _conn()
    try:
        # 模拟"讲到第 3 页、第 2 个讲解单元"（页序 0 起算 → slide=2）
        repo.update_session(conn, first["session_id"], position_slide=2,
                            position_unit=1, elapsed_sec=18.0)
    finally:
        conn.close()
    client.post(f"/api/v1/sessions/{first['session_id']}/end")

    # 结束这一刻必须已经把最终位置写进 progress（不能只靠运行期节流写）
    conn = _conn()
    try:
        prog = repo.get_progress(conn, client.get("/api/v1/profiles").json()["active_id"],
                                 unit_id)
    finally:
        conn.close()
    assert prog is not None
    assert prog["last_slide"] == 2
    assert prog["last_unit"] == 1
    assert prog["study_sec"] == 18.0

    again = _session_via_api(client, unit_id)
    assert again["session_id"] != first["session_id"]     # 新会话（上一节已结束）
    assert again["resumed"] is True                       # 但接着上次的进度
    assert again["position"] == {"slide": 2, "unit_index": 1, "from_progress": True}
    # 课程详情里的单元进度也要对得上（前端"继续"按钮读的就是它）
    detail = client.get("/api/v1/courses/c1").json()
    unit_view = detail["chapters"][0]["units"][0]
    assert unit_view["progress"]["last_slide"] == 2


def test_session_other_profile_starts_at_zero(client):
    """换档案 → 同一个单元从第 1 页开始（进度互不干扰）。"""
    unit_id = _seed_course(1)[0]
    first = _session_via_api(client, unit_id)
    conn = _conn()
    try:
        repo.update_session(conn, first["session_id"], position_slide=2, position_unit=1)
    finally:
        conn.close()
    client.post(f"/api/v1/sessions/{first['session_id']}/end")

    other = client.post("/api/v1/profiles", json={"name": "另一个人"}).json()
    fresh = client.post("/api/v1/sessions",
                        json={"unit_id": unit_id, "profile_id": other["id"]}).json()
    assert fresh["resumed"] is False
    assert fresh["position"] == {"slide": 0, "unit_index": 0, "from_progress": False}


def test_session_end_and_recap(client):
    unit_id = _seed_course(1)[0]
    sid = _session_via_api(client, unit_id)["session_id"]
    ended = client.post(f"/api/v1/sessions/{sid}/end").json()
    assert ended["ok"] is True
    recap = ended["recap"]
    assert set(recap) == {"elapsed_sec", "slides_covered", "units_given",
                          "qa_count", "topics"}
    assert client.get(f"/api/v1/sessions/{sid}").json()["status"] == "ended"
    # recap 可重放：刷新小结页读到的还是同一份
    again = client.get(f"/api/v1/sessions/{sid}/recap").json()
    assert again == recap


def test_session_recap_persisted_as_event(client):
    unit_id = _seed_course(1)[0]
    sid = _session_via_api(client, unit_id)["session_id"]
    client.post(f"/api/v1/sessions/{sid}/end")
    conn = _conn()
    try:
        event = repo.latest_event(conn, sid, "recap")
    finally:
        conn.close()
    assert event is not None and event["payload"]["slides_covered"] >= 0


def test_session_delete_leaves_no_orphans(client):
    unit_id = _seed_course(1)[0]
    sid = _session_via_api(client, unit_id)["session_id"]
    conn = _conn()
    try:
        repo.add_event(conn, sid, "ping", {})
        repo.delete_session(conn, sid)
        assert conn.execute("SELECT COUNT(*) AS n FROM events WHERE session_id=?",
                            (sid,)).fetchone()["n"] == 0
        assert conn.execute("SELECT COUNT(*) AS n FROM sessions WHERE id=?",
                            (sid,)).fetchone()["n"] == 0
    finally:
        conn.close()


def test_session_validation(client):
    no_unit = client.post("/api/v1/sessions", json={})
    assert no_unit.status_code == 400
    assert "unit_id" in no_unit.json()["message"]
    assert client.post("/api/v1/sessions", json={"unit_id": "nope"}).status_code == 404
    unit_id = _seed_course(1)[0]
    assert client.post("/api/v1/sessions",
                       json={"unit_id": unit_id, "profile_id": "nope"}).status_code == 404


# ===========================================================================
# /dev 开关
# ===========================================================================

def test_dev_is_404_by_default(client):
    unit_id = _seed_course(1)[0]
    sid = _session_via_api(client, unit_id)["session_id"]
    assert client.get(f"/dev/units/{unit_id}/plan").status_code == 404
    assert client.get(f"/dev/sessions/{sid}/events").status_code == 404


def test_dev_on_serves_events(client, monkeypatch):
    monkeypatch.setenv("CLASSROOM_DEV", "1")
    unit_id = _seed_course(1)[0]
    sid = _session_via_api(client, unit_id)["session_id"]
    conn = _conn()
    try:
        repo.add_event(conn, sid, "ping", {"n": 1})
    finally:
        conn.close()
    body = client.get(f"/dev/sessions/{sid}/events").json()
    assert body["count"] == 1
    assert body["items"][0]["kind"] == "ping"
    assert body["items"][0]["payload"] == {"n": 1}
    # plan.json 不在这台机器的课件目录里 → 404 而不是 500
    assert client.get(f"/dev/units/{unit_id}/plan").status_code == 404


# ===========================================================================
# 事件上限
# ===========================================================================

def test_events_capped_per_session():
    conn = _conn()
    try:
        sid = repo.create_session(conn, "u1", "p1")["id"]
        for i in range(repo.EVENTS_KEEP + 120):
            repo.add_event(conn, sid, "tick", {"i": i})
        assert repo.count_events(conn, sid) == repo.EVENTS_KEEP
        newest = repo.list_events(conn, sid, 1)[0]
        assert newest["payload"]["i"] == repo.EVENTS_KEEP + 119   # 留的是最近的
    finally:
        conn.close()


# ===========================================================================
# WS（真运行器 + stub 课件）
# ===========================================================================

def test_ws_hello_page_and_segments(client, monkeypatch):
    _install_stub_engine(monkeypatch)
    unit_id = _seed_course(1)[0]
    sid = _session_via_api(client, unit_id)["session_id"]

    with client.websocket_connect(f"/ws/classroom/{sid}") as ws:
        frames = _pump(ws, until=lambda f: f["kind"] == "hello", limit=5)
        hello = frames[-1]
        assert hello["kind"] == "hello"
        assert set(hello) >= {"kind", "unit", "profile", "position"}
        assert hello["unit"]["id"] == unit_id
        assert hello["position"] == {"slide": 0, "unit_index": 0}
        assert set(hello["profile"]) == {"id", "name"}

        frames = _pump(ws, until=lambda f: f["kind"] == "page", limit=5)
        page = frames[-1]
        assert page["slide_index"] == 0
        assert page["slide_count"] == 4
        assert page["slide"]["mode"] == "simple"
        assert page["media_base"] == ""             # 没有真实课件 → 没有媒体前缀

        frames = _pump(ws, until=lambda f: f["kind"] == "segment_start", limit=20)
        seg = frames[-1]
        assert seg["seq"] >= 1
        assert seg["text"]
        assert seg["seg_kind"] == "unit"
        frames = _pump(ws, until=lambda f: f["kind"] == "segment_end", limit=20)
        assert frames[-1]["seq"] == seg["seq"]
        # 没有 TTS → 绝不该收到 audio 帧
        assert "audio" not in [f["kind"] for f in frames]

def test_ws_pause_and_resume(client, monkeypatch):
    _install_stub_engine(monkeypatch)
    unit_id = _seed_course(1)[0]
    sid = _session_via_api(client, unit_id)["session_id"]

    with client.websocket_connect(f"/ws/classroom/{sid}") as ws:
        _pump(ws, until=lambda f: f["kind"] == "segment_start", limit=20)
        ws.send_json({"t": "cmd", "cmd": "pause"})
        frames = _drain_until_state(ws, "已暂停")
        assert any(f["kind"] == "notify" and f["level"] == "info" for f in frames)
        conn = _conn()
        try:
            assert repo.get_session(conn, sid)["status"] == "paused"
        finally:
            conn.close()

        ws.send_json({"t": "cmd", "cmd": "resume"})
        frames = _drain_until_state(ws, "继续")
        assert any(f["kind"] == "state" for f in frames)
        conn = _conn()
        try:
            assert repo.get_session(conn, sid)["status"] == "active"
        finally:
            conn.close()


def test_ws_resume_replays_from_same_segment(client, monkeypatch):
    # speech_timeout 拉长到 1s：暂停一定落在"第一句正播到一半"的窗口里，
    # 续播必须重发这一句（而不是跳到下一句，也不是从头开始）。
    _install_stub_engine(monkeypatch, speech_timeout=1.0)
    unit_id = _seed_course(1)[0]
    sid = _session_via_api(client, unit_id)["session_id"]

    with client.websocket_connect(f"/ws/classroom/{sid}") as ws:
        # 同步点放在 segment_end：这一刻服务端一定正处在"等这句话播完"的 1s 窗口里，
        # 于是紧接着的暂停**必然**打断这一句 —— 续播必须重发它。
        frames = _pump(ws, until=lambda f: f["kind"] == "segment_end", limit=20)
        interrupted = [f for f in frames if f["kind"] == "segment_start"][-1]
        ws.send_json({"t": "cmd", "cmd": "pause"})
        _drain_until_state(ws, "已暂停")
        ws.send_json({"t": "cmd", "cmd": "resume"})
        drained = _drain_until_state(ws, "继续")
        # 续播帧可能夹在"继续"和它的 state 之间（两条协程并发推送），先在这批里找
        replay = next((f for f in drained if f["kind"] == "segment_start"
                       and f["seq"] > interrupted["seq"]), None)
        if replay is None:
            frames = _pump(ws, until=lambda f: f["kind"] == "segment_start"
                           and f["seq"] > interrupted["seq"], limit=30)
            replay = frames[-1]
        assert replay["seq"] > interrupted["seq"]      # 新序号（前端不会把它当旧段）
        assert replay["text"] == interrupted["text"]    # 同一句在重发（不是跳到下一句）


def test_ws_speech_done_drives_advance(client, monkeypatch):
    """推进必须由客户端的 {t:"speech_done",seq} 驱动，而且序号要绑定。

    这里把 speech_timeout 拉到 10s：如果 speech_done 没被认，"下一段"要等 10 秒才来，
    用例会先超时 —— 于是"用时 < 5s"本身就证明是回执在推进课堂。
    先发一个错配的 seq（必须被忽略：上一段迟到的 done 落到这一段上会让整节课复读）。
    """
    _install_stub_engine(monkeypatch, speech_timeout=10.0)
    unit_id = _seed_course(1)[0]
    sid = _session_via_api(client, unit_id)["session_id"]

    with client.websocket_connect(f"/ws/classroom/{sid}") as ws:
        frames = _pump(ws, until=lambda f: f["kind"] == "segment_end", limit=20)
        seq = frames[-1]["seq"]
        started = time.monotonic()
        ws.send_json({"t": "speech_done", "seq": seq + 99})     # 错配 → 忽略
        ws.send_json({"t": "speech_done", "seq": seq})          # 正确 → 推进
        frames = _pump(ws, until=lambda f: f["kind"] == "segment_start" and f["seq"] > seq,
                       limit=20)
        elapsed = time.monotonic() - started
        assert frames[-1]["seq"] == seq + 1
        assert elapsed < 5.0, f"推进应来自 speech_done（实测 {elapsed:.2f}s ≈ 超时兜底）"


def test_ws_stray_speech_done_is_ignored(client, monkeypatch):
    """错配的 speech_done 不许结束等待（旧 Runner 的 test_protocol 语义之一）。

    为什么这条重要：一次错配（上一段的 done 落在这一段上）就会让"等播完"落到错误的
    段落上，状态机提前推进 —— 学生看到的是老师同一句话被跳过、后面全错位。
    做法：speech_timeout 10s，收到 segment_end 后先回一个 seq+99（必须被忽略），
    断言 1 秒内**一帧都没有**；再回正确的 seq，必须立刻推进到 seq+1。
    """
    _install_stub_engine(monkeypatch, speech_timeout=10.0)
    unit_id = _seed_course(1)[0]
    sid = _session_via_api(client, unit_id)["session_id"]

    with client.websocket_connect(f"/ws/classroom/{sid}") as ws:
        frames = _frame_reader(ws)
        seg = _wait_kind(frames, "segment_end")
        assert seg is not None, "没等到 segment_end"
        seq = seg["seq"]

        # 段末的旁路帧（字幕等）先读干净，再开始"安静期"观察
        while (frame := _next_frame(frames, 0.3)) is not None:
            assert frame["kind"] != "segment_start", "同一段还没播完就推进了"

        ws.send_json({"t": "speech_done", "seq": seq + 99})      # 错配
        assert _next_frame(frames, 1.0) is None,             "错配的 speech_done 不许推进课堂（也不该产生任何帧）"

        ws.send_json({"t": "speech_done", "seq": seq})           # 正确
        nxt = _wait_kind(frames, "segment_start", timeout=5.0)
        assert nxt is not None, "正确的 speech_done 没有推进课堂"
        assert nxt["seq"] == seq + 1

    # 错配这件事必须留下痕迹（/dev 面板与日志靠它排查"为什么卡住"）
    conn = _conn()
    try:
        kinds = [e["kind"] for e in repo.list_events(conn, sid)]
    finally:
        conn.close()
    assert "stray_speech_done" in kinds


def test_ws_early_speech_done_advances_immediately(client, monkeypatch):
    """提前回执也要被尊重（旧 Runner 的 test_protocol 语义之二）。

    前端可以在 `segment_start` 一到就回 `speech_done`（不等音频真的播完）。
    speech_timeout 给 10s：若回执被忽略，下一段要等 10 秒才出现 ——
    实测必须在 2 秒内到达，否则就是"等超时兜底"而不是"被尊重"。
    """
    _install_stub_engine(monkeypatch, speech_timeout=10.0)
    unit_id = _seed_course(1)[0]
    sid = _session_via_api(client, unit_id)["session_id"]

    with client.websocket_connect(f"/ws/classroom/{sid}") as ws:
        frames = _frame_reader(ws)
        seg = _wait_kind(frames, "segment_start")
        assert seg is not None, "没等到 segment_start"
        started = time.monotonic()
        ws.send_json({"t": "speech_done", "seq": seg["seq"]})    # 收到就开始播 → 立刻回执
        nxt = _wait_kind(frames, "segment_start", timeout=5.0)   # 中间夹的 segment_end 跳过
        elapsed = time.monotonic() - started
        assert nxt is not None, "提前回执没有推进课堂"
        assert nxt["seq"] == seg["seq"] + 1
        assert elapsed < 2.0, f"提前回执必须立即推进（实测 {elapsed:.2f}s ≈ 等超时）"


def test_ws_next_sends_next_page(client, monkeypatch):
    _install_stub_engine(monkeypatch)
    unit_id = _seed_course(1)[0]
    sid = _session_via_api(client, unit_id)["session_id"]

    with client.websocket_connect(f"/ws/classroom/{sid}") as ws:
        _pump(ws, until=lambda f: f["kind"] == "page" and f["slide_index"] == 0, limit=5)
        ws.send_json({"t": "cmd", "cmd": "next"})
        frames = _pump(ws, until=lambda f: f["kind"] == "page" and f["slide_index"] == 1,
                       limit=40)
        assert frames[-1]["slide_index"] == 1


def test_ws_bad_messages_get_notify(client, monkeypatch):
    _install_stub_engine(monkeypatch)
    unit_id = _seed_course(1)[0]
    sid = _session_via_api(client, unit_id)["session_id"]

    with client.websocket_connect(f"/ws/classroom/{sid}") as ws:
        _pump(ws, until=lambda f: f["kind"] == "hello", limit=5)
        ws.send_text("这不是 JSON")
        frames = _pump(ws, until=lambda f: f["kind"] == "notify" and "不是 JSON" in f["text"],
                       limit=20)
        assert frames[-1]["level"] == "warn"
        ws.send_json({"t": "rate", "value": 3.0})
        frames = _pump(ws, until=lambda f: f["kind"] == "notify" and "语速" in f["text"],
                       limit=20)
        assert frames[-1]["level"] == "warn"


def test_ws_unknown_session_reports_and_closes(client):
    with client.websocket_connect("/ws/classroom/s-does-not-exist") as ws:
        frame = ws.receive_json()
        assert frame["kind"] == "notify"
        assert frame["level"] == "error"


# ===========================================================================
# 运行器（不经 WS：暂停真的停住、落库、recap）
# ===========================================================================

async def _new_runner(*, slides: int = 4, speech_timeout: float = 0.5):
    from studio.runner import SessionRunner

    conn = _conn()
    unit = {"id": "u-run", "course_id": "c-run", "title": "运行器单元",
            "slides": slides, "pptx_path": ""}
    session = repo.create_session(conn, "u-run", "p-run")
    collector = _Collector()
    runner = SessionRunner(
        session=session, unit=unit, profile={"id": "p-run", "name": "测试"}, send=collector,
        conn=conn, settings=_mini_settings(), tts=None, course=_stub_course(slides),
        llm=_StubLLM(), speech_timeout=speech_timeout)
    return runner, collector, conn


async def test_runner_speaks_and_pauses_completely():
    runner, collector, conn = await _new_runner()
    try:
        await runner.start()
        assert await _await_until(lambda: "segment_start" in collector.kinds())
        await runner.pause()
        settled = len(collector.frames)
        await asyncio.sleep(0.25)
        # 暂停期间一帧都不许再推（连 notify/state 也不行）
        assert len(collector.frames) == settled, collector.kinds()[settled:]
        await runner.resume()
        assert await _await_until(lambda: len(collector.frames) > settled)
    finally:
        await runner.close()
        conn.close()


async def test_runner_persists_progress_and_position():
    runner, collector, conn = await _new_runner()
    try:
        await runner.start()
        assert await _await_until(lambda: "segment_start" in collector.kinds())
        await asyncio.sleep(0.2)
        row = repo.get_session(conn, runner.session_id)
        assert row["status"] == "active"
        assert row["elapsed_sec"] >= 0
        prog = repo.get_progress(conn, "p-run", "u-run")
        assert prog is not None and prog["updated_at"]
    finally:
        await runner.close()
        conn.close()


async def test_runner_finish_returns_recap_and_ends():
    runner, collector, conn = await _new_runner()
    try:
        await runner.start()
        assert await _await_until(lambda: "segment_start" in collector.kinds())
        recap = await runner.finish()
        assert set(recap) >= {"elapsed_sec", "slides_covered", "units_given",
                              "qa_count", "topics"}
        assert repo.get_session(conn, runner.session_id)["status"] == "ended"
        assert any(f["kind"] == "recap" for f in collector.frames)
        assert repo.latest_event(conn, runner.session_id, "recap") is not None
        # F1.2：结束这节课时，最终位置必须已经落在 progress 里（续学的唯一依据）
        prog = repo.get_progress(conn, "p-run", "u-run")
        assert prog is not None
        assert prog["last_slide"] == runner.state.slide_index
    finally:
        await runner.close()
        conn.close()


async def test_runner_close_marks_paused_for_resume():
    runner, collector, conn = await _new_runner()
    try:
        await runner.start()
        assert await _await_until(lambda: "segment_start" in collector.kinds())
        sid = runner.session_id
        await runner.close()
        row = repo.get_session(conn, sid)
        assert row["status"] == "paused"          # 断线 = 暂停，可续学
        assert repo.resumable_session(conn, "p-run", "u-run") is not None
    finally:
        conn.close()


async def test_runner_never_emits_audio_without_tts():
    runner, collector, conn = await _new_runner()
    try:
        await runner.start()
        assert await _await_until(lambda: "segment_start" in collector.kinds())
        assert "audio" not in collector.kinds()
    finally:
        await runner.close()
        conn.close()
