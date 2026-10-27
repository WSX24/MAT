# -*- coding: utf-8 -*-
"""studio/outline.py：单元大纲（形状钉死 §4，数据来自 deck.json + plan.json）。

为什么这些用例值得存在：大纲页是"这单元讲什么"的唯一入口，它裂一次就等于这个
单元点不开。所以这里钉三件事：
  - 页数与 deck.json 一致、totals 加得对（前端标题栏直接用这几个数）；
  - plan.json 没有时**降级成规则大纲**而不是报错（刚导入还没备课是常态）；
  - preview 截断到 ~60 字、kp 缺失给空数组（契约里没有 null）。

夹具直接写 deck.json / plan.json（大纲只读这两份产物），另外留一条真备料 →
大纲的整链用例，免得"两份产物拼不在一起"这种错被手写 JSON 掩盖过去。
"""

from __future__ import annotations

import json
import sys
from dataclasses import replace
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from config import ProviderProfile, load_settings    # noqa: E402
from deck import store as deck_store                 # noqa: E402
from studio import materialize, outline               # noqa: E402


# ===========================================================================
# 夹具：直接写产物（大纲只读 deck.json / plan.json）
# ===========================================================================

def _deck_json(pages: list[tuple[str, list[str]]]) -> dict:
    """(页标题, [正文段落…]) → deck.json 的最小形状。

    只保留大纲真正读的那几个键（index / title / elements[*].role / text.paras）：
    夹具跟着**契约**走，而不是跟着某个版本的解析器输出走。
    """
    slides = []
    for i, (title, paras) in enumerate(pages):
        elements = [{"element_id": f"s{i + 1:02d}-sh01", "role": "title",
                     "text": {"paras": [{"pid": f"s{i + 1:02d}-sh01-p01", "text": title}]}}]
        if paras:
            elements.append({"element_id": f"s{i + 1:02d}-sh02", "role": "body",
                             "text": {"paras": [{"pid": f"s{i + 1:02d}-sh02-p{j + 1:02d}",
                                                 "text": t} for j, t in enumerate(paras)]}})
        slides.append({"index": i, "title": title, "elements": elements})
    return {"deck_id": "demo", "title": "演示课件", "slides": slides}


def _plan(units: list[dict]) -> dict:
    return {"deck_id": "demo", "source": "rule", "units": units}


def _write(base: Path, name: str, data: dict) -> Path:
    base.mkdir(parents=True, exist_ok=True)
    path = base / name
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    return path


def _pptx(path: Path, slides: int = 3) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i in range(slides):
        page = prs.slides.add_slide(prs.slide_layouts[5])
        page.shapes.title.text = f"第 {i + 1} 页：运动的描述"
        box = page.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2))
        box.text_frame.text = f"第 {i + 1} 页正文，质点与参考系先讲清楚。"
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def _keyless_settings():
    base = load_settings()
    keyless = ProviderProfile(name="deepseek", base_url="https://api.deepseek.com/v1",
                              api_key="", model="deepseek-chat")
    return replace(base, provider="deepseek", providers={"deepseek": keyless})


# ===========================================================================
# 用例
# ===========================================================================

def test_outline_follows_deck_pages_and_plan_units(tmp_path) -> None:
    """页数与 deck.json 一致，单元按 slide 挂到各自那页，totals 加得对。"""
    _write(tmp_path, "deck.json", _deck_json([
        ("第一页", ["质点是一个有质量的点。"]),
        ("第二页", ["参考系要先选定。"]),
        ("第三页", ["时刻与时间不是一回事。"]),
    ]))
    _write(tmp_path, "plan.json", _plan([
        {"slide": 0, "text": "质点是一个有质量的点，先把模型讲清楚。", "kp": ["质点"]},
        {"slide": 0, "text": "为什么要有质点这个模型。", "est_sec": 12.0, "kp": []},
        {"slide": 1, "text": "参考系要先选定，不然运动说不清。", "kp": ["参考系"]},
    ]))

    out = outline.build_outline(tmp_path)

    assert set(out) == {"totals", "pages"}, "正常路径的形状严格等于 §4 契约"
    assert [p["index"] for p in out["pages"]] == [0, 1, 2]
    assert [p["title"] for p in out["pages"]] == ["第一页", "第二页", "第三页"]
    assert out["totals"]["slides"] == 3
    assert out["totals"]["units"] == 3
    assert [len(p["units"]) for p in out["pages"]] == [2, 1, 0], "备课刻意跳过的页老实留空"
    first = out["pages"][0]["units"]
    assert first[0]["kp"] == ["质点"]
    assert first[1]["est_sec"] == 12.0, "plan 里写了 est_sec 就用它"
    assert first[0]["est_sec"] == round(len("质点是一个有质量的点，先把模型讲清楚。") / 5.0, 1)
    assert out["totals"]["est_sec"] > 0
    assert round(sum(u["est_sec"] for p in out["pages"] for u in p["units"]), 1) \
        == out["totals"]["est_sec"]


def test_without_plan_every_page_gets_one_rule_unit(tmp_path) -> None:
    """没有 plan.json：一页一个单元，preview 取正文首段（不是标题），别报错。"""
    _write(tmp_path, "deck.json", _deck_json([
        ("第一页", ["正文首段在这里。", "第二段不该被选中。"]),
        ("第二页", ["这一页也有正文。"]),
        ("第三页", ["第三页正文。"]),
    ]))

    out = outline.build_outline(tmp_path)

    assert set(out) == {"totals", "pages"}, "降级路径不该额外多出 warnings"
    assert [len(p["units"]) for p in out["pages"]] == [1, 1, 1]
    assert out["pages"][0]["units"][0]["preview"] == "正文首段在这里。"
    assert out["pages"][0]["units"][0]["kp"] == [], "没有 kp 就给空数组，不给 null"
    assert all(u["est_sec"] > 0 for p in out["pages"] for u in p["units"])
    assert set(out["totals"]) == {"slides", "units", "est_sec"}, "totals 只有契约里这三个键"
    assert out["totals"]["est_sec"] == round(
        sum(u["est_sec"] for p in out["pages"] for u in p["units"]), 1), "totals 要对得上明细"


def test_preview_is_clipped_to_one_line(tmp_path) -> None:
    """preview 压成一行并截到 ~60 字（投屏一行 / 手机一屏）。"""
    long_text = "质点" * 200
    _write(tmp_path, "deck.json", _deck_json([("只有一页", ["正文。"])]))
    _write(tmp_path, "plan.json", _plan([
        {"slide": 0, "text": f"开头\n{long_text}", "kp": []},
    ]))

    view = outline.build_outline(tmp_path)["pages"][0]["units"][0]

    assert "\n" not in view["preview"]
    assert len(view["preview"]) <= outline.PREVIEW_CHARS + 1
    assert view["preview"].endswith("…"), "截断了才加省略号"


def test_kp_missing_or_blank_is_an_empty_array(tmp_path) -> None:
    """kp 缺失 / null / 全是空白串 → 空数组（契约里没有 null）。"""
    _write(tmp_path, "deck.json", _deck_json([("页", ["正文。"])]))
    _write(tmp_path, "plan.json", _plan([
        {"slide": 0, "text": "没有 kp 的单元。"},
        {"slide": 0, "text": "kp 是 null 的单元。", "kp": None},
        {"slide": 0, "text": "kp 里有空串的单元。", "kp": ["", "质点", "  "]},
    ]))

    units = outline.build_outline(tmp_path)["pages"][0]["units"]

    assert [u["kp"] for u in units] == [[], [], ["质点"]]


def test_missing_deck_json_is_an_empty_outline_not_a_crash(tmp_path) -> None:
    """deck.json 不在：空大纲 + 一句 warnings（不抛，页面还能打开）。"""
    out = outline.build_outline(tmp_path / "还没有备料")

    assert out["totals"] == {"slides": 0, "units": 0, "est_sec": 0.0}
    assert out["pages"] == []
    assert any("prepare" in w for w in out["warnings"]), out["warnings"]


def test_broken_plan_json_falls_back_to_rule_outline(tmp_path) -> None:
    """plan.json 是坏的：退回规则大纲，并在 warnings 里说清是哪份文件。"""
    _write(tmp_path, "deck.json", _deck_json([("页一", ["正文一。"]), ("页二", ["正文二。"])]))
    (tmp_path / "plan.json").write_text("{ 这不是 json", encoding="utf-8")

    out = outline.build_outline(tmp_path)

    assert [len(p["units"]) for p in out["pages"]] == [1, 1], "坏计划 → 规则大纲"
    assert any("plan.json 读不动" in w for w in out["warnings"]), out["warnings"]


def test_empty_plan_keeps_pages_but_says_there_are_no_units(tmp_path) -> None:
    """plan.json 在、但一个单元都没有：页数照给，单元留空并说明（不假装备过课）。"""
    _write(tmp_path, "deck.json", _deck_json([("页一", ["正文一。"]), ("页二", ["正文二。"])]))
    _write(tmp_path, "plan.json", _plan([]))

    out = outline.build_outline(tmp_path)

    assert out["totals"]["slides"] == 2
    assert out["totals"]["units"] == 0
    assert [p["units"] for p in out["pages"]] == [[], []]
    assert any("没有讲解单元" in w for w in out["warnings"]), out["warnings"]


def test_est_sec_matches_script_gen_rule(tmp_path) -> None:
    """est_sec 的"5 字/秒"必须与 script_gen.PlanUnit.est_sec() 同一口径。"""
    from script_gen import PlanUnit

    text = "字" * 50
    assert PlanUnit(slide=0, text=text).est_sec() == 10.0

    _write(tmp_path, "deck.json", _deck_json([("页", ["正文。"])]))
    _write(tmp_path, "plan.json", _plan([{"slide": 0, "text": text}]))

    unit = outline.build_outline(tmp_path)["pages"][0]["units"][0]

    assert unit["est_sec"] == 10.0, "两份产物算出来的时长必须一样，否则列表与大纲对不上"


def test_outline_of_a_really_prepared_unit(tmp_path) -> None:
    """整链：真备料 → 大纲。页数与 deck.json 一致，且每页都有可讲的东西。"""
    src = _pptx(tmp_path / "第1单元-运动的描述.pptx", slides=3)
    decks = tmp_path / "decks"
    prepared = materialize.prepare_unit(src, deck_root=decks, settings=_keyless_settings())
    assert prepared["ok"] is True, prepared["warnings"]

    out = outline.build_outline(decks / deck_store.make_deck_id(src.name))

    assert out["totals"]["slides"] == prepared["slides"] == 3
    assert out["totals"]["units"] >= 3, "规则备课每页至少一个单元"
    assert out["totals"]["est_sec"] > 0
    assert all(p["units"] for p in out["pages"]), "规则备课时没有哪页是空的"
