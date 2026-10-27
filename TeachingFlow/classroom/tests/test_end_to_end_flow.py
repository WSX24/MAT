"""端到端主流程（设计文档 §0 主线 / §8.1-§8.3）：一份真 pptx 从进库到续学。

这条用例只证明一件事：**主线是通的**。所以它不替任何一层做替身 ——
真 pptx（python-pptx 现场造 3 页）→ 真导入 → 真备课（规则兜底）→ 真课程库
→ 真大纲 → 真运行器（WS 逐帧收，按协议回 speech_done）→ 真小结 → 真续学。

离线口径：配置给一个连不上的 stub 服务商（备课走规则兜底、讲解照读课件）、TTS 关掉，
全程不联网、不调模型、不花额度。

隔离口径：库、runs、state（导入的 pptx）、**decks/**（解析产物）全部落 tmp_path。
最后一项最容易被忽略 —— 导入会真的解析 pptx 并写 decks/<deck_id>/，落到仓库里就会
污染 decks/，连带把"按固定 deck_id 选真实课件"的那几条示例用例带偏（见
tests/fixtures.py 的 REAL_DECK_PREFIXES）。
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

CLASSROOM = Path(__file__).resolve().parents[1]
if str(CLASSROOM) not in sys.path:              # 允许从任意 cwd 直接跑本文件
    sys.path.insert(0, str(CLASSROOM))

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

#: 讲的页数上限保护：WS 帧是**会合式**的（服务端多推一帧、测试不读，双方就各等各的），
#: 所以每个"等某帧"的地方都要有次数上界，卡死要变成红的断言，不是挂住的用例。
FRAME_LIMIT = 200


# ===========================================================================
# 夹具
# ===========================================================================

def _offline_settings():
    """连不上的 stub 服务商：任何模型调用都秒失败 → 备课走规则兜底。"""
    from config import ProviderProfile, Settings

    return Settings(
        provider="stub",
        providers={"stub": ProviderProfile(name="stub", base_url="http://127.0.0.1:9/v1",
                                           api_key="", model="stub")},
        temperature=0.2, timeout=2.0,
        tts_enabled=False, tts_voice="", tts_rate="+0%",
        probe_window=0.05, max_clarify_asks=1, max_rephrase=1, unit_char_warn=200,
    )


@pytest.fixture()
def env(tmp_path, monkeypatch):
    """临时库 + 临时 runs + 临时 state + 临时 decks，全部关在 tmp_path 里。"""
    monkeypatch.setenv("CLASSROOM_DB", str(tmp_path / "classroom.sqlite"))
    monkeypatch.delenv("CLASSROOM_DEV", raising=False)
    from app import deps
    from deck import store as deck_store

    monkeypatch.setattr(deps, "STATE_ROOT", tmp_path / "state", raising=False)
    monkeypatch.setattr(deps, "RUNS_ROOT", tmp_path / "runs", raising=False)
    monkeypatch.setattr(deps, "DIST_ROOT", tmp_path / "dist", raising=False)
    (tmp_path / "runs").mkdir()
    (tmp_path / "decks").mkdir()
    monkeypatch.setattr(deck_store, "DECKS_DIR", tmp_path / "decks", raising=False)
    monkeypatch.setattr(deps, "get_settings", _offline_settings)
    deps.reset_settings_cache()
    yield tmp_path
    deps.reset_settings_cache()


@pytest.fixture()
def client(env):
    from app import create_app

    with TestClient(create_app()) as c:
        yield c


def _make_pptx(path: Path, slides: int = 3) -> bytes:
    """3 页课件：每页一个标题 + 一段正文（规则备课会把它们切成讲解单元）。"""
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    for i in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title = slide.shapes.add_textbox(Emu(600000), Emu(400000), Emu(7000000), Emu(700000))
        title.text_frame.text = f"第 {i + 1} 页 标题"
        body = slide.shapes.add_textbox(Emu(600000), Emu(1300000), Emu(7000000), Emu(2200000))
        body.text_frame.text = (f"第 {i + 1} 页第一句要讲的话。"
                               f"第 {i + 1} 页第二句，补一点细节。")
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path.read_bytes()


# ===========================================================================
# WS 小工具
# ===========================================================================

def _read_until(ws, predicate, *, limit: int = FRAME_LIMIT, what: str = "目标帧") -> list[dict]:
    """读到 predicate(frame, 已收到的全部帧) 为真；读满次数就失败（不挂住）。"""
    frames: list[dict] = []
    for _ in range(limit):
        frame = ws.receive_json()
        frames.append(frame)
        if predicate(frame, frames):
            return frames
    kinds = [f.get("kind") for f in frames]
    pytest.fail(f"等了 {limit} 帧也没等到{what}；收到的 kind 序列：{kinds}")


def _play(ws, segments: int, *, limit: int = FRAME_LIMIT) -> list[dict]:
    """按协议上课：每收到一个 segment_start 就回 speech_done（"这一段播完了"）。

    这正是真前端的做法（设计文档 §4 C→S `speech_done{seq}`）：不回它，状态机会停在
    同一段上等，学生看到的是老师反复说同一句话。
    """
    frames: list[dict] = []
    done = 0
    for _ in range(limit):
        frame = ws.receive_json()
        frames.append(frame)
        if frame["kind"] == "segment_start":
            ws.send_json({"t": "speech_done", "seq": frame.get("seq")})
            done += 1
            if done >= segments:
                return frames
    kinds = [f.get("kind") for f in frames]
    pytest.fail(f"讲了 {done}/{segments} 段就停了；收到的 kind 序列：{kinds}")


def _send_cmd(ws, cmd: str, *, marker: str = "", expect_kind: str = "state",
              limit: int = FRAME_LIMIT) -> list[dict]:
    """发一条控制命令，读到它的回执帧。

    pause/resume 的回执是「notify + state」两帧连发：`state.phase` 说的是状态机当前
    的相位（LECTURING），**不是** "paused" —— 暂停这件事由 notify 的文案表达。
    所以断言点必须落在"先看到那条 notify，再看到它的 state"上。
    """
    ws.send_json({"t": "cmd", "cmd": cmd})

    seen_marker = [not marker]

    def hit(frame: dict, _all: list[dict]) -> bool:
        if marker and frame.get("kind") == "notify" and marker in (frame.get("text") or ""):
            seen_marker[0] = True
            return False
        return seen_marker[0] and frame.get("kind") == expect_kind

    return _read_until(ws, hit, limit=limit, what=f"cmd={cmd} 的 {expect_kind} 回执")


# ===========================================================================
# 主线
# ===========================================================================

def test_import_course_learn_and_resume(client, env):
    """导入 → 课程库 → 详情 → 大纲 → 开课 → 上课（暂停/继续/翻页）→ 小结 → 续学。"""
    # ---- ① 导入一份真 pptx ------------------------------------------------
    data = _make_pptx(env / "src" / "三页课件.pptx")
    resp = client.post("/api/v1/courses/import",
                       files={"file": ("三页课件.pptx", data, PPTX_MIME)})
    assert resp.status_code == 200, resp.text
    imported = resp.json()
    course_id, unit_id = imported["course_id"], imported["unit_id"]
    assert imported["slides"] == 3, f"页数必须来自真解析：{imported}"

    # 解析产物落在临时 decks/ 里，绝不能写进仓库
    from deck import store as deck_store

    assert str(deck_store.DECKS_DIR).startswith(str(env)), "解析产物必须落在临时目录"

    # ---- ② 课程库列表看得到 ------------------------------------------------
    courses = client.get("/api/v1/courses").json()
    row = next((c for c in courses if c["id"] == course_id), None)
    assert row is not None, f"导入的课程没进列表：{courses}"
    assert row["source"] == "import"
    assert row["unit_count"] == 1

    # ---- ③ 课程详情看得到单元 ----------------------------------------------
    detail = client.get(f"/api/v1/courses/{course_id}").json()
    units = [u for ch in detail["chapters"] for u in ch["units"]]
    assert [u["id"] for u in units] == [unit_id]
    assert units[0]["slides"] == 3

    # ---- ④ 大纲有页 --------------------------------------------------------
    outline = client.get(f"/api/v1/units/{unit_id}/outline").json()
    assert outline["pages"], f"大纲不能是空的：{outline}"
    assert outline["totals"]["slides"] == 3
    unit_detail = client.get(f"/api/v1/units/{unit_id}").json()
    assert unit_detail["ready"] is True, f"导入时已经备过一次课：{unit_detail}"

    # ---- ⑤ 开课：WS 收到 hello --------------------------------------------
    created = client.post("/api/v1/sessions", json={"unit_id": unit_id}).json()
    session_id = created["session_id"]
    assert created["resumed"] is False, "第一次开课不该是续学"
    assert created["position"]["slide"] == 0

    with client.websocket_connect(f"/ws/classroom/{session_id}") as ws:
        frames = _read_until(ws, lambda f, _a: f.get("kind") == "hello", what="hello 帧")
        hello = frames[-1]
        assert hello["unit"]["id"] == unit_id
        assert hello["position"]["slide"] == 0
        assert "profile" in hello and "unit" in hello

        # 第一段：等 segment_start 就回 speech_done
        played = _play(ws, 2)
        assert any(f["kind"] == "segment_start" for f in played)
        assert any(f["kind"] == "segment_end" for f in played), "播完必须收到 segment_end"

        # ---- ⑥ 控制条：暂停 / 继续 / 下一页 都不报错 ----------------------
        paused = _send_cmd(ws, "pause", marker="已暂停")
        assert not [f for f in paused if f.get("level") == "error"], f"暂停不该报错：{paused}"

        resumed = _send_cmd(ws, "resume", marker="继续")
        assert not [f for f in resumed if f.get("level") == "error"], f"继续不该报错：{resumed}"

        # 继续之后必须真的接着讲（否则"继续"只是改了个标志位）
        _play(ws, 1)
        _send_cmd(ws, "next", expect_kind="page")
        _play(ws, 1)
        errors = []
        # 收尾：把还可能在路上的一两帧读掉，顺便确认没有 error 级 notify
        for _ in range(6):
            try:
                frame = ws.receive_json()
            except Exception:               # 通道里没帧了 —— 正常，不该为它红
                break
            if frame.get("kind") == "notify" and frame.get("level") == "error":
                errors.append(frame)
        assert errors == [], f"整节课不该出现 error 级 notify：{errors}"

    # ---- ⑦ 结束：拿到 recap ------------------------------------------------
    ended = client.post(f"/api/v1/sessions/{session_id}/end")
    assert ended.status_code == 200, ended.text
    recap = ended.json()["recap"]
    assert set(recap) >= {"elapsed_sec", "slides_covered", "units_given", "qa_count", "topics"}
    assert recap["slides_covered"] >= 2, f"翻过页就该覆盖到第 2 页：{recap}"
    assert client.get(f"/api/v1/sessions/{session_id}/recap").json() == recap, \
        "recap 必须可重放（同一份）"

    # ---- ⑧ 续学：同一个单元再开一次，从上次位置接着上 --------------------
    again = client.post("/api/v1/sessions", json={"unit_id": unit_id}).json()
    assert again["session_id"] != session_id, "上一轮已结束，该开新会话"
    assert again["resumed"] is True, f"学到一半再来必须算续学：{again}"
    assert again["position"]["slide"] != 0, f"续学位置不许回到第 1 页：{again}"
    assert again["position"]["from_progress"] is True

    # 换档案：进度互不干扰（设计文档 §8.3）
    other = client.post("/api/v1/profiles", json={"name": "另一个学生"}).json()
    fresh = client.post("/api/v1/sessions",
                        json={"unit_id": unit_id, "profile_id": other["id"]}).json()
    assert fresh["resumed"] is False, "新档案必须从第一页开始"
    assert fresh["position"]["slide"] == 0
