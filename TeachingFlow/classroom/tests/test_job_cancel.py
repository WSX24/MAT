"""P2-3：被取消/超时的作业不能永远停在 running。

以前 _run_job 在 CancelledError 里只写一行日志就 re-raise，status 还是 running；
而 _running_for 只认 running、_prune 又拒绝删 running —— 这份源之后所有作业都被
"正在预处理"拒掉。这份测试钉三件事：状态不再 running、error 是人话、同一份源
能立刻再起一个新作业（单飞解锁）。
"""

from __future__ import annotations

import asyncio
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import preprocess                                # noqa: E402
from deck import store as deck_store             # noqa: E402


@pytest.fixture()
def decks(tmp_path, monkeypatch) -> Path:
    root = tmp_path / "decks"
    root.mkdir()
    monkeypatch.setattr(deck_store, "DECKS_DIR", root)
    return root


def make_pptx(path: Path, slides: int = 3) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i in range(slides):
        page = prs.slides.add_slide(prs.slide_layouts[5])
        page.shapes.title.text = f"第 {i + 1} 页：取消测试"
        box = page.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2))
        box.text_frame.text = f"第 {i + 1} 页正文。"
    prs.save(str(path))
    return path


def test_wait_job_sees_a_failed_record_after_cancel(tmp_path, decks, monkeypatch) -> None:
    """同步外壳，方便读断言：起作业 → cancel → 状态归一。"""
    src = make_pptx(tmp_path / "取消掉.pptx")
    started = asyncio.Event()

    async def _slow_ingest(rec, src_path, deck_id):
        started.set()
        await asyncio.sleep(30)          # 真的不会睡满：下面 0.2s 量级就取消
        raise AssertionError("取消之后不该继续往下跑")

    monkeypatch.setattr(preprocess, "_stage_ingest", _slow_ingest)

    async def _main() -> tuple[dict, dict]:
        job = await preprocess.start_job(src.name, action="ingest", source_path=src)
        await asyncio.wait_for(started.wait(), timeout=5)   # 确保真的进了阶段
        snap = await preprocess.cancel_job(job["id"], reason="测试取消")
        return job, snap

    job, snap = asyncio.run(_main())

    assert snap["status"] != "running", "取消之后不许还是 running（单飞会被永久占住）"
    assert snap["status"] == "failed", snap
    assert "取消" in (snap["error"] or ""), snap["error"]
    assert snap["finished_at"], "取消也要有结束时间"
    assert "取消" in "\n".join(snap["log_tail"]), snap["log_tail"][-3:]
    # 记录表里读到的是同一份结论（前端轮询走的就是 get_job）
    assert preprocess.get_job(job["id"])["status"] == "failed"


def test_source_can_be_reused_after_cancel(tmp_path, decks, monkeypatch) -> None:
    """单飞解锁：同一份源在旧作业被取消之后必须能再起一个。"""
    src = make_pptx(tmp_path / "解锁.pptx")
    started = asyncio.Event()

    async def _slow_ingest(rec, src_path, deck_id):
        started.set()
        await asyncio.sleep(30)

    monkeypatch.setattr(preprocess, "_stage_ingest", _slow_ingest)

    async def _main() -> tuple[dict, dict]:
        job = await preprocess.start_job(src.name, action="ingest", source_path=src)
        await asyncio.wait_for(started.wait(), timeout=5)
        await preprocess.cancel_job(job["id"], reason="测试取消")
        again = await preprocess.start_job(src.name, action="ingest", source_path=src)
        # 收尾：第二个也取消掉，别把 30s 睡眠留到测试结束之后
        await preprocess.cancel_job(again["id"], reason="测试清理")
        return job, again

    job, again = asyncio.run(_main())

    assert again["id"] != job["id"]
    assert again["status"] != "failed" or "正在预处理" not in (again["error"] or ""), again
    assert "正在预处理" not in (again["error"] or ""), "同一份源被单飞拒掉了"


def test_cancel_of_a_job_that_never_started_still_unlocks(tmp_path, decks, monkeypatch) -> None:
    """任务还没被调度就被取消：协程体一次都没跑，finally 也不会执行。

    这条是最阴的一档 —— 光靠 try/finally 归一不出来，所以 cancel_job 自己兜一次。
    """
    src = make_pptx(tmp_path / "还没开始.pptx")

    async def _slow_ingest(rec, src_path, deck_id):
        await asyncio.sleep(30)

    monkeypatch.setattr(preprocess, "_stage_ingest", _slow_ingest)

    async def _main() -> dict:
        job = await preprocess.start_job(src.name, action="ingest", source_path=src)
        return await preprocess.cancel_job(job["id"], reason="立刻取消")

    snap = asyncio.run(_main())
    assert snap["status"] == "failed", snap
    assert "取消" in (snap["error"] or ""), snap["error"]
    assert snap["finished_at"]
def test_bridge_timeout_cancels_the_job_and_unlocks_the_source(tmp_path, decks, monkeypatch) -> None:
    """bridge 的超时分支：光报一句"超时"不够，那份作业会把单飞一直占着。

    这里用 0.05s 的超时 + 一个 5 秒的阶段函数：bridge 一轮轮询之后就会超时，
    取消并等它收尾，随后同一份源必须能再起作业。
    """
    import bridge

    src = make_pptx(tmp_path / "会超时.pptx")

    async def _slow_ingest(rec, src_path, deck_id):
        await asyncio.sleep(5)

    monkeypatch.setattr(preprocess, "_stage_ingest", _slow_ingest)

    res = asyncio.run(bridge.preprocess_pptx(src, action="ingest", timeout=0.05))
    assert res["ok"] is False
    assert "超时" in res["error"], res["error"]
    assert "取消" in res["error"], "超时后必须真的取消了作业，不能只报一句"
    assert "failed" in res["stages"], res["stages"]

    async def _restart() -> dict:
        job = await preprocess.start_job(src.name, action="ingest", source_path=src)
        await preprocess.cancel_job(job["id"], reason="测试清理")
        return job

    job = asyncio.run(_restart())
    assert job.get("error") is None, job
