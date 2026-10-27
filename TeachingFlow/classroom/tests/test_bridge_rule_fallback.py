"""P2-2：没有可用模型时，bridge 与直连入口要落到**同一种产物**。

直连入口（run.py plan / /api/preprocess/start action=plan）在模型不可用时按页回退
成规则计划（script_gen.plan_rule，本地、秒级、照读课件），作业报 done；而 bridge
以前直接跳过备课只做解析。同一份 pptx，走出片链路和走网页得到的不是同一节课。

现在 bridge 也走规则兜底：ingest → plan_rule → 预置路线 → 落盘，ok=True 且
degraded=True。这份测试钉的就是"plan.json 真的写出来了，而且第一单元的 note
写着规则兜底"。
"""

from __future__ import annotations

import asyncio
import json
import sys
from dataclasses import replace
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import bridge                                    # noqa: E402
from config import ProviderProfile, load_settings    # noqa: E402
from deck import store as deck_store             # noqa: E402


def make_pptx(path: Path, slides: int = 3) -> Path:
    """现场生成小课件，不依赖仓库里那份真实课件。"""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i in range(slides):
        page = prs.slides.add_slide(prs.slide_layouts[5])       # Title Only
        page.shapes.title.text = f"第 {i + 1} 页：规则兜底"
        box = page.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2))
        box.text_frame.text = f"第 {i + 1} 页正文，规则备课要有字可讲。"
    prs.save(str(path))
    return path


@pytest.fixture()
def decks(tmp_path, monkeypatch) -> Path:
    root = tmp_path / "decks"
    root.mkdir()
    monkeypatch.setattr(deck_store, "DECKS_DIR", root)
    return root


def _keyless_settings():
    """没有可用模型的一份配置：远程服务商但 api_key 为空（连探活都不会发出去）。"""
    base = load_settings()
    keyless = ProviderProfile(name="deepseek", base_url="https://api.deepseek.com/v1",
                              api_key="", model="deepseek-chat")
    return replace(base, provider="deepseek", providers={"deepseek": keyless})


def test_full_without_model_writes_a_rule_plan(tmp_path, decks) -> None:
    src = make_pptx(tmp_path / "无模型也要有计划.pptx")
    res = asyncio.run(bridge.preprocess_pptx(src, action="full", settings=_keyless_settings()))

    deck_id = deck_store.make_deck_id(src.name)
    plan_path = decks / deck_id / "plan.json"

    assert res["ok"] is True, res["error"]
    assert res["degraded"] is True, res["note"]
    assert res["action"] == "full", res["action"]
    assert "规则兜底" in res["note"], res["note"]
    assert "action='plan'" in res["note"], "note 要告诉人配好 key 之后怎么换正式备课"

    assert plan_path.is_file(), "无模型时也要落 plan.json（否则两个入口产物不一致）"
    raw = json.loads(plan_path.read_text(encoding="utf-8"))
    assert len(raw["units"]) > 0, "规则兜底也要切出单元"
    assert "规则兜底" in raw["units"][0]["note"], raw["units"][0]
    # 与 run.py plan 同一套落盘口径：顶层带 routes 键（没有岔路时是空数组）
    assert isinstance(raw.get("routes"), list), raw.keys()

    assert res["stats"]["has_plan"] is True
    assert res["stats"]["units"] == len(raw["units"])
    assert "plan" in res["stages"] and "routes" in res["stages"], res["stages"]


def test_rule_fallback_never_calls_the_model(tmp_path, decks, monkeypatch) -> None:
    """不许"先试模型再回退"：没有可用模型时 plan_llm 一次都不该被调用。"""
    import script_gen

    calls: list[int] = []

    async def _boom(*args, **kwargs):
        calls.append(1)
        raise AssertionError("没有可用模型时不该去调模型（会白等逐页超时）")

    monkeypatch.setattr(script_gen, "plan_llm", _boom)

    src = make_pptx(tmp_path / "不调模型.pptx")
    res = asyncio.run(bridge.preprocess_pptx(src, action="full", settings=_keyless_settings()))

    assert res["ok"] is True, res["error"]
    assert calls == [], "走到 LLM 备课了"
