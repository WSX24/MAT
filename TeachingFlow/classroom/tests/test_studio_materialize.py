# -*- coding: utf-8 -*-
"""studio/materialize.py：备料（解析 + 备课 + 落盘），钉四件事。

1. 真的产出了 deck.json + plan.json，而且规模（页数/时长）是从盘上数出来的；
2. 没有可用模型时走规则兜底，并如实说 degraded —— 不假装备过课；
3. 第二次调用不再解析（幂等）：**bridge 一次都不该被调**，这是"进入单元"快慢的分水岭；
4. 失败不抛，warnings 里是人话。

用例全部用 tmp_path 与现场生成的小课件，deck 落点用 deck_root 指到 tmp：
不联网、不调模型、绝不写 classroom/decks 真目录。
"""

from __future__ import annotations

import sys
from dataclasses import replace
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import bridge                                          # noqa: E402
from config import ProviderProfile, load_settings      # noqa: E402
from deck import store as deck_store                   # noqa: E402
from studio import materialize                         # noqa: E402

#: 契约字段一个都不能少、也不许多（W1 直接把它塞进 /units/{id}/prepare 的响应）
_CONTRACT = {"ok", "deck_id", "deck_ready", "plan_ready", "plan_source",
             "degraded", "slides", "est_sec", "warnings"}


def _pptx(path: Path, slides: int = 3) -> Path:
    """造一份 3 页小课件：每页有标题与一句正文（备课要"有字可讲"）。"""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i in range(slides):
        page = prs.slides.add_slide(prs.slide_layouts[5])       # Title Only
        page.shapes.title.text = f"第 {i + 1} 页：运动的描述"
        box = page.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2))
        box.text_frame.text = (f"第 {i + 1} 页正文，质点、参考系、时刻与时间，"
                               "这几个词以后每天都会用到。")
    prs.save(str(path))
    return path


def _keyless_settings():
    """一份没有可用模型的配置：远程服务商但 api_key 为空（连探活都不会发出去）。

    显式传它，是为了让用例与这台机器上有没有 .env key 无关 —— 否则"没模型时
    走规则兜底"这条会随环境飘。
    """
    base = load_settings()
    keyless = ProviderProfile(name="deepseek", base_url="https://api.deepseek.com/v1",
                              api_key="", model="deepseek-chat")
    return replace(base, provider="deepseek", providers={"deepseek": keyless})


@pytest.fixture()
def decks(tmp_path) -> Path:
    """deck 落点（既有 classroom/decks/ 的替身）。"""
    return tmp_path / "decks"


# ===========================================================================
# 用例
# ===========================================================================

def test_prepare_unit_writes_deck_and_rule_plan(tmp_path, decks) -> None:
    """真备一次：deck.json 与 plan.json 都落盘，页数/时长从盘上数得出来。"""
    src = _pptx(tmp_path / "第1单元-运动的描述.pptx", slides=3)

    res = materialize.prepare_unit(src, deck_root=decks, settings=_keyless_settings())

    deck_id = deck_store.make_deck_id(src.name)
    assert set(res) == _CONTRACT, res.keys()
    assert res["ok"] is True, res["warnings"]
    assert res["deck_id"] == deck_id
    assert res["deck_ready"] is True and res["plan_ready"] is True
    assert (decks / deck_id / "deck.json").is_file()
    assert (decks / deck_id / "plan.json").is_file()
    assert res["plan_source"] == "rule", "没有可用模型 → 规则兜底"
    assert res["degraded"] is True
    assert res["slides"] == 3
    assert res["est_sec"] > 0, "规则兜底也要讲得出时长（按讲稿字数估）"
    assert any("规则兜底" in w for w in res["warnings"]), res["warnings"]


def test_second_call_reuses_and_never_parses(tmp_path, decks, monkeypatch) -> None:
    """幂等：second call 直接读盘返回，bridge 一次都不该被调。"""
    src = _pptx(tmp_path / "第1单元-运动的描述.pptx")
    first = materialize.prepare_unit(src, deck_root=decks, settings=_keyless_settings())
    assert first["ok"] is True, first["warnings"]

    def _boom(*args, **kwargs):
        raise AssertionError("已经备好的单元不该再解析一遍（几十秒 + 模型钱）")

    monkeypatch.setattr(bridge, "preprocess_pptx_sync", _boom)

    second = materialize.prepare_unit(src, deck_root=decks, settings=_keyless_settings())

    assert set(second) == _CONTRACT, second.keys()
    assert second["ok"] is True
    assert second["deck_ready"] is True and second["plan_ready"] is True
    assert (second["slides"], second["est_sec"]) == (first["slides"], first["est_sec"])
    assert second["plan_source"] == first["plan_source"] == "rule"
    assert any("直接复用" in w for w in second["warnings"]), second["warnings"]


def test_force_parses_again(tmp_path, decks, monkeypatch) -> None:
    """force=True 才重跑一次（换了一版课件、想重新备课）。"""
    src = _pptx(tmp_path / "第1单元-运动的描述.pptx")
    materialize.prepare_unit(src, deck_root=decks, settings=_keyless_settings())

    calls: list[str] = []
    real = bridge.preprocess_pptx_sync

    def _count(src_arg, **kw):
        calls.append(str(kw.get("action")))
        return real(src_arg, **kw)

    monkeypatch.setattr(bridge, "preprocess_pptx_sync", _count)

    res = materialize.prepare_unit(src, deck_root=decks, settings=_keyless_settings(),
                                  force=True)

    assert calls == ["full"], "force 走的就是完整链路（解析 + 备课 + 落盘）"
    assert res["ok"] is True and res["deck_ready"] is True


def test_missing_pptx_speaks_instead_of_raising(tmp_path, decks) -> None:
    """课件不在盘上：ok=False + 人话，不抛。"""
    res = materialize.prepare_unit(tmp_path / "没有这份.pptx", deck_root=decks)

    assert set(res) == _CONTRACT, res.keys()
    assert res["ok"] is False
    assert res["deck_ready"] is False and res["plan_ready"] is False
    assert any("不在盘上" in w for w in res["warnings"]), res["warnings"]


def test_non_pptx_is_refused_with_a_human_reason(tmp_path, decks) -> None:
    """别的格式在进解析器之前就拦下来（解析器只会给一句看不懂的 zip 报错）。"""
    doc = tmp_path / "讲稿.md"
    doc.write_text("# 这不是课件", encoding="utf-8")

    res = materialize.prepare_unit(doc, deck_root=decks)

    assert res["ok"] is False
    assert any(".pptx" in w for w in res["warnings"]), res["warnings"]


def test_broken_pptx_fails_with_a_human_reason(tmp_path, decks) -> None:
    """坏 pptx：ok=False，warnings 里是"另存一次再传"那种人话，不是 ZipError。"""
    bad = tmp_path / "坏掉的课件.pptx"
    bad.write_bytes("这不是 zip".encode("utf-8"))

    res = materialize.prepare_unit(bad, deck_root=decks, settings=_keyless_settings())

    assert res["ok"] is False
    assert res["deck_ready"] is False
    assert any("打不开" in w or "另存" in w for w in res["warnings"]), res["warnings"]


def test_deck_root_is_restored_after_the_call(tmp_path, decks) -> None:
    """deck_root 只是这一次调用的开关，退出（含异常）必须还原。"""
    before = deck_store.DECKS_DIR
    src = _pptx(tmp_path / "第1单元-运动的描述.pptx")

    materialize.prepare_unit(src, deck_root=decks, settings=_keyless_settings())

    assert deck_store.DECKS_DIR == before, "把全局落点留在 tmp 里会毒到后面的调用"
